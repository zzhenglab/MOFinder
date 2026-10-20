"""Meaningful offline checks of denominators, units, alignment, and provenance."""

import csv
from contextlib import redirect_stdout
import io
import itertools
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from mofinder.evaluation import positive as p


def reference_row(record_id="P001-R001", doi="10.example/a", name="MOF-A", **fields):
    return {**dict.fromkeys(p.FIELDS, ""), "record_id": record_id,
            "paper_id": record_id.split("-")[0], "doi": doi, "mof_name": name, **fields}


class NormalizationTests(unittest.TestCase):
    def test_both_delimiters_and_explicit_readable_absence(self):
        for reference in ("not reported | water", "__ABSENT__ ||| water"):
            self.assertTrue(p.field_label("", reference, "solvent_main")["exact_correct"])
            self.assertTrue(p.field_label("H2O", reference, "solvent_main")["exact_correct"])
            self.assertTrue(p.field_label("", reference, "solvent_main")["optional_reference"])
            self.assertFalse(p.field_label("DMF", reference, "solvent_main")["exact_correct"])
        self.assertEqual(p.display_alternatives("DMF ||| water"), "DMF | water")
        self.assertEqual(p.alternatives("DMF; water"), ["DMF; water"])

    def test_explicit_unit_equivalence_and_no_inferred_mass_conversion(self):
        self.assertTrue(p.value_matches("720 min", "12", "time_h"))
        self.assertTrue(p.value_matches("0.01 g", "10 mg", "metal_1_amount_text"))
        self.assertTrue(p.value_matches("0.1 mmol, 0.0297 g", "29.7 mg", "metal_1_amount_text"))
        self.assertFalse(p.value_matches("10 mg", "0.1 mmol", "metal_1_amount_text"))
        self.assertFalse(p.value_matches("12 h then 4 h", "12", "time_h"))
        self.assertFalse(p.value_matches("48-72 h", "48", "time_h"))
        self.assertFalse(p.value_matches("0.1 mol/L", "0.1 mol", "metal_1_amount_text"))

    def test_names_keep_hydrates_and_mixtures_distinct(self):
        self.assertTrue(p.value_matches("DMF", "N,N-dimethylformamide", "solvent_main"))
        self.assertFalse(p.value_matches("zinc nitrate", "zinc nitrate hexahydrate", "metal_1"))
        self.assertFalse(p.value_matches("water", "water (in DES)", "solvent_main"))
        self.assertTrue(p.value_matches("H2O", "water (in DES)", "solvent_main", ["water"]))
        self.assertFalse(p.value_matches("DMF", "DMF; water", "solvent_main"))
        self.assertTrue(p.value_matches("DEF", "DMF ||| DEF", "solvent_main"))
        self.assertFalse(p.is_absent("Na"))
        self.assertFalse(p.is_absent(0))

    def test_doi_wrappers_and_numeric_identifier(self):
        self.assertEqual(p.doi_key(" HTTPS://doi.org/10.1234%2FABC "), "10.1234/abc")
        self.assertTrue(p.value_matches("123.0", "123", "crystal_code"))
        self.assertTrue(p.value_matches("393.15 K", "120", "temperature_c"))
        self.assertFalse(p.value_matches("120", "122", "temperature_c"))

    def test_projection_keeps_text_amount_priority_and_explicit_fallbacks(self):
        projected = p.projected_row({
            "metal_1": "zinc chloride", "metal_1_abbr": "ZnCl2",
            "metal_1_amount_text": "10 mg", "metal_1_amount_value": 99,
            "metal_1_amount_unit": "mg", "linker_1_abbr": "BDC",
            "linker_1_amount_value": 0.2, "linker_1_amount_unit": "mmol",
            "solvent_main_abbr": "DMF", "solvent_main_amount_text": "2 mL",
            "solvent_main_ml": 9, "solvent_secondary_abbr": "H2O",
            "solvent_secondary_ml": 0.5, "temperature_c": 0,
            "temperature_c_text": "room temperature", "time_text": "48-72 h",
        })
        self.assertEqual(projected["metal_1"], "zinc chloride")
        self.assertEqual(projected["metal_1_amount_text"], "10 mg")
        self.assertEqual(projected["linker_1"], "BDC")
        self.assertEqual(projected["linker_1_amount_text"], "0.2 mmol")
        self.assertEqual(projected["solvent_main_amount_text"], "2 mL")
        self.assertEqual(projected["solvent_secondary_amount_text"], "0.5 mL")
        self.assertEqual(projected["temperature_c"], "0")
        self.assertEqual(projected["time_h"], "48-72 h")


class DenominatorTests(unittest.TestCase):
    def test_wrong_value_is_fp_and_fn_but_one_cell(self):
        label = p.field_label("99", "12", "time_h")
        summary = p.summarize_fields([label])
        self.assertEqual((summary["TP"], summary["FP"], summary["FN"]), (0, 1, 1))
        self.assertEqual(summary["assessable_cells"], 1)
        self.assertEqual(summary["exact_cell_accuracy"], 0)
        self.assertEqual(summary["presence_accuracy"], 1)
        self.assertIsNone(summary["presence_balanced_accuracy"])

    def test_empty_unmatched_slots_do_not_create_true_negatives(self):
        for alignment in ("missed", "extra"):
            label = p.field_label("", "", "time_h", alignment=alignment)
            self.assertFalse(label["assessable"])
            self.assertEqual(label["TN"], 0)
        label = p.field_label("", "", "time_h")
        self.assertEqual(label["TN"], 1)

    def test_unknown_reference_is_not_a_missing_value(self):
        label = p.field_label("99", "__UNRESOLVED__", "time_h")
        self.assertFalse(label["assessable"])
        self.assertEqual(label["FP"], 0)
        self.assertEqual(p.summarize_fields([label])["unresolved_cells"], 1)

    def test_union_with_approved_absence_has_no_single_presence_label(self):
        expected = "__ABSENT__ ||| 12"
        blank = p.field_label("", expected, "time_h")
        correct = p.field_label("720 min", expected, "time_h")
        wrong = p.field_label("99", expected, "time_h")
        missed = p.field_label("", expected, "time_h", alignment="missed")
        self.assertTrue(blank["exact_correct"])
        self.assertEqual(blank["TN"], 1)
        self.assertEqual(correct["TP"], 1)
        self.assertEqual((wrong["FP"], wrong["FN"]), (1, 0))
        self.assertFalse(missed["assessable"])
        summary = p.summarize_fields([blank, correct, wrong, missed])
        self.assertEqual(summary["optional_reference_cells"], 4)
        self.assertEqual(summary["presence_assessable_cells"], 0)
        self.assertIsNone(summary["presence_accuracy"])
        self.assertIsNone(summary["required_value_recall"])

    def test_record_tn_and_dependent_metrics_are_undefined(self):
        result = p.confusion_metrics(2, 1, None, 1)
        for key in ("TN", "accuracy", "balanced_accuracy"):
            self.assertIsNone(result[key])
        self.assertAlmostEqual(result["F1"], 2 / 3)


class PooledFieldTests(unittest.TestCase):
    @staticmethod
    def label(field, predicted, expected, alignment="matched"):
        return {"field": field, "alignment": alignment, "doi": "10.example/a",
                **p.field_label(predicted, expected, field, alignment=alignment)}

    def test_groups_partition_all_slots_without_reuse(self):
        members = [field for fields in p.FIELD_GROUPS.values() for field in fields]
        self.assertEqual(len(p.FIELD_GROUPS), 14)
        self.assertCountEqual(members, p.FIELDS)
        self.assertEqual(len(members), len(set(members)))
        self.assertEqual(p.FIELD_GROUPS["metal_name"], ("metal_1", "metal_2", "metal_3"))
        self.assertEqual(p.FIELD_GROUPS["linker_name"], ("linker_1", "linker_2", "linker_3"))
        self.assertEqual(p.FIELD_GROUPS["solvent_name"], ("solvent_main", "solvent_secondary"))

    def test_unequal_slot_denominators_pool_counts_and_keep_amounts_separate(self):
        labels = [
            self.label("metal_1", "ZnCl2", "ZnCl2"),
            *[self.label("metal_2", "CuCl2", "") for _ in range(9)],
            self.label("metal_3", "", ""),
            self.label("metal_1_amount_text", "2 mg", "1 mg"),
            self.label("metal_2_amount_text", "", ""),
            self.label("metal_3_amount_text", "", "1 mg", "missed"),
        ]
        metrics = {row["field"]: row for row in p.summarize_field_groups(labels)
                   if row["scope"] == "all_records"}
        names, amounts = metrics["metal_name"], metrics["metal_amount"]
        self.assertEqual(tuple(names[key] for key in ("TP", "FP", "TN", "FN")), (1, 9, 1, 0))
        self.assertAlmostEqual(names["precision"], 0.1)
        self.assertNotEqual(names["precision"], (1.0 + 0.0) / 2)
        self.assertEqual(tuple(amounts[key] for key in ("TP", "FP", "TN", "FN")), (0, 1, 1, 2))
        self.assertEqual(amounts["assessable_cells"], 3)
        self.assertAlmostEqual(amounts["exact_cell_accuracy"], 1 / 3)

    def test_pooling_preserves_every_count_and_exact_cell_denominator_in_both_scopes(self):
        labels = [self.label(field, "", "") for field in p.FIELDS]
        labels += [
            self.label("metal_3", "ZnCl2", "ZnCl2"),
            self.label("metal_3_amount_text", "2 mg", "1 mg"),
            self.label("linker_2", "BDC", "__UNREVIEWED__"),
            self.label("solvent_main", "", "not reported | water"),
            self.label("solvent_secondary", "DMF", "not reported | water"),
            self.label("solvent_main", "", "not reported | water", "missed"),
            self.label("time_h", "", "12", "missed"),
            self.label("linker_3", "BDC", "", "extra"),
            self.label("linker_2", "", "", "extra"),
        ]
        metrics = p.summarize_field_groups(labels)
        reference = {"papers": [{"paper_id": "P001", "doi": "10.example/a", "included": True}]}
        summary, _, _ = p.summarize([], labels, reference, {})
        additive_counts = (
            "TP", "FP", "TN", "FN", "assessable_cells", "unresolved_cells", "wrong_values",
            "optional_reference_cells", "optional_assessable_cells", "presence_assessable_cells",
            "required_value_TP", "required_value_FN",
            "presence_TP", "presence_FP", "presence_TN", "presence_FN",
        )
        for scope in ("all_records", "matched_records"):
            with self.subTest(scope=scope):
                selected = [row for row in labels
                            if scope == "all_records" or row["alignment"] == "matched"]
                expected = p.summarize_fields(selected)
                self.assertEqual(summary["field_evaluation"][scope]["micro"], expected)
                pooled = [row for row in metrics if row["scope"] == scope]
                self.assertEqual(len(pooled), 14)
                for key in additive_counts:
                    self.assertEqual(sum(row[key] for row in pooled), expected[key], key)
                correct_cells = sum(row["assessable"] and row["exact_correct"] is True
                                    for row in selected)
                pooled_correct_cells = sum(
                    row["exact_cell_accuracy"] * row["assessable_cells"]
                    for row in pooled if row["exact_cell_accuracy"] is not None)
                self.assertAlmostEqual(pooled_correct_cells, correct_cells)
                self.assertAlmostEqual(
                    pooled_correct_cells / sum(row["assessable_cells"] for row in pooled),
                    expected["exact_cell_accuracy"])
                self.assertEqual(expected["optional_reference_cells"],
                                 3 if scope == "all_records" else 2)
                self.assertEqual(expected["optional_assessable_cells"], 2)

    def test_macro_averages_defined_pooled_types_equally(self):
        labels = [
            self.label("metal_1", "ZnCl2", "ZnCl2"),
            *[self.label("metal_2", "CuCl2", "") for _ in range(9)],
            self.label("linker_3", "BDC", "BDC"),
            self.label("time_h", "99", "12"),
        ]
        reference = {"papers": [{"paper_id": "P001", "doi": "10.example/a", "included": True}]}
        summary, metrics, _ = p.summarize([], labels, reference, {})
        fields = summary["field_evaluation"]["all_records"]
        self.assertAlmostEqual(fields["macro_fields"]["precision"], (0.1 + 1 + 0) / 3)
        self.assertEqual(fields["macro_fields_defined_counts"]["precision"], 3)
        self.assertAlmostEqual(fields["micro"]["precision"], 2 / 12)
        self.assertNotEqual(fields["macro_fields"]["precision"], (1 + 0 + 1 + 0) / 4)
        self.assertEqual(len(metrics), 28)
        self.assertIn("14 pooled field types", summary["definitions"]["macro_fields"])


class RoleSlotTests(unittest.TestCase):
    def test_tied_unknown_entities_have_content_based_order(self):
        expected = reference_row()
        first = reference_row(modulator_1="acetic acid", modulator_1_amount_text="10 mL")
        second = reference_row(modulator_2="acetic acid", modulator_2_amount_text="10 mL")
        aligned_first, _ = p.align_role_slots(first, expected)
        aligned_second, _ = p.align_role_slots(second, expected)
        for field in p.FIELDS:
            self.assertEqual(aligned_first[field], aligned_second[field])
        self.assertEqual(p.pair_score(first, expected), p.pair_score(second, expected))

    def test_optional_reference_does_not_make_unknown_solvent_order_sensitive(self):
        expected = reference_row(solvent_main="DMF | not reported")
        first = reference_row(solvent_main="water", solvent_main_amount_text="10 mL")
        second = reference_row(solvent_secondary="water", solvent_secondary_amount_text="10 mL")
        aligned_first, _ = p.align_role_slots(first, expected)
        aligned_second, _ = p.align_role_slots(second, expected)
        for field in p.FIELDS:
            self.assertEqual(aligned_first[field], aligned_second[field])
        self.assertEqual(aligned_first["solvent_main"], "water")
        self.assertEqual(p.pair_score(first, expected), p.pair_score(second, expected))

    def test_all_role_permutations_include_third_metal_and_third_linker(self):
        fields = {}
        for group in p.ROLE_GROUPS:
            for position, slot in enumerate(group, 1):
                fields[slot] = f"{slot}-identity"
                fields[slot + "_amount_text"] = f"{position} mmol"
        expected = reference_row(**fields)
        for permutations in itertools.product(
                *(list(itertools.permutations(group)) for group in p.ROLE_GROUPS)):
            candidate = dict(expected)
            for group, permutation in zip(p.ROLE_GROUPS, permutations):
                for target, source in zip(group, permutation):
                    candidate[target] = expected[source]
                    candidate[target + "_amount_text"] = expected[source + "_amount_text"]
            aligned, _ = p.align_role_slots(candidate, expected)
            for field in fields:
                self.assertTrue(p.value_matches(aligned[field], expected[field], field),
                                (permutations, field))
            self.assertEqual(p.pair_score(candidate, expected), p.pair_score(expected, expected))

    def test_solvent_names_and_amounts_can_swap_together(self):
        expected = reference_row(solvent_main="DMF", solvent_main_amount_text="5 mL",
                                 solvent_secondary="water", solvent_secondary_amount_text="1 mL")
        candidate = reference_row(solvent_main="H2O", solvent_main_amount_text="1 mL",
                                  solvent_secondary="DMF", solvent_secondary_amount_text="5 mL")
        aligned, mapping = p.align_role_slots(candidate, expected)
        self.assertEqual(mapping["solvent_main"], "solvent_secondary")
        for field in ("solvent_main", "solvent_main_amount_text",
                      "solvent_secondary", "solvent_secondary_amount_text"):
            self.assertTrue(p.value_matches(aligned[field], expected[field], field))
        self.assertEqual(p.pair_score(candidate, expected), p.pair_score(expected, expected))

    def test_metal_names_and_amounts_can_swap_together(self):
        expected = reference_row(metal_1="ZnCl2", metal_1_amount_text="10 mg",
                                 metal_2="CuCl2", metal_2_amount_text="20 mg")
        candidate = reference_row(metal_1="CuCl2", metal_1_amount_text="20 mg",
                                  metal_2="ZnCl2", metal_2_amount_text="10 mg")
        aligned, mapping = p.align_role_slots(candidate, expected)
        self.assertEqual(mapping["metal_1"], "metal_2")
        self.assertEqual(aligned["metal_1_amount_text"], "10 mg")

    def test_amounts_cannot_be_swapped_independently_to_manufacture_correctness(self):
        expected = reference_row(metal_1="ZnCl2", metal_1_amount_text="10 mg",
                                 metal_2="CuCl2", metal_2_amount_text="20 mg")
        candidate = reference_row(metal_1="ZnCl2", metal_1_amount_text="20 mg",
                                  metal_2="CuCl2", metal_2_amount_text="10 mg")
        aligned, mapping = p.align_role_slots(candidate, expected)
        self.assertEqual(mapping["metal_1"], "metal_1")
        self.assertFalse(p.value_matches(aligned["metal_1_amount_text"],
                                         expected["metal_1_amount_text"], "metal_1_amount_text"))
        self.assertFalse(p.value_matches(aligned["metal_2_amount_text"],
                                         expected["metal_2_amount_text"], "metal_2_amount_text"))

    def test_one_duplicate_identity_cannot_be_reused_for_two_reference_slots(self):
        expected = reference_row(metal_1="ZnCl2", metal_1_amount_text="10 mg",
                                 metal_2="ZnCl2", metal_2_amount_text="20 mg")
        candidate = reference_row(metal_1="ZnCl2", metal_1_amount_text="10 mg")
        aligned, mapping = p.align_role_slots(candidate, expected)
        self.assertEqual(len(set(mapping[k] for k in ("metal_1", "metal_2", "metal_3"))), 3)
        self.assertEqual(sum(not p.is_absent(aligned[k]) for k in ("metal_1", "metal_2")), 1)

    def test_merged_name_amount_union_cannot_create_unobserved_pairs(self):
        expected = reference_row(solvent_main="DMF ||| water",
                                 solvent_main_amount_text="10 mL ||| 20 mL")
        pairs = {"solvent_main": [
            {"name": "water", "amount_text": "10 mL"},
            {"name": "DMF", "amount_text": "20 mL", "amount_aliases": ["0.020 L"]},
        ]}
        amount, aliases, status = p.amount_reference(
            "solvent_main", "H2O", expected, role_pairs=pairs)
        self.assertEqual(status, "paired_variant")
        self.assertFalse(p.value_matches("20 mL", amount, "solvent_main_amount_text", aliases))
        self.assertTrue(p.value_matches("10 mL", amount, "solvent_main_amount_text", aliases))

    def test_paired_amount_compares_all_shared_dimensions(self):
        expected = reference_row(metal_1="HfPiv", metal_1_amount_text="403 mg, 0.167 mmol")
        pairs = {"metal_1": [{"name": "HfPiv", "amount_text": "403 mg, 0.167 mmol"}]}
        amount, aliases, _ = p.amount_reference("metal_1", "HfPiv", expected, role_pairs=pairs)
        for value in ("403 mg", "0.167 mmol", "0.403 g, 0.167 mmol"):
            self.assertTrue(p.value_matches(value, amount, "metal_1_amount_text", aliases))
        self.assertFalse(p.value_matches("403 mg, 1 mmol", amount, "metal_1_amount_text", aliases))

    def test_explicit_literal_variant_has_no_inferred_volume(self):
        expected = reference_row(solvent_secondary="__ABSENT__ ||| 100ml")
        pairs = {"solvent_secondary": [
            {"name": "", "amount_text": ""},
            {"name": "100ml", "amount_text": ""},
        ]}
        self.assertTrue(p.field_label("100ml", expected["solvent_secondary"],
                                      "solvent_secondary")["exact_correct"])
        self.assertTrue(p.field_label("", expected["solvent_secondary"],
                                      "solvent_secondary")["exact_correct"])
        self.assertFalse(p.value_matches("water", expected["solvent_secondary"], "solvent_secondary"))
        amount, aliases, _ = p.amount_reference("solvent_secondary", "100ml", expected,
                                                role_pairs=pairs)
        self.assertTrue(p.field_label("", amount, "solvent_secondary_amount_text",
                                      aliases)["exact_correct"])
        self.assertFalse(p.field_label("100 mL", amount, "solvent_secondary_amount_text",
                                       aliases)["exact_correct"])


class RecordRuleTests(unittest.TestCase):
    @staticmethod
    def composite_rule():
        return {"composite_projections": [{
            "from_slots": ["modulator_1", "modulator_2"],
            "source_names": ["triethylamine + water"],
            "source_amounts": ["TEA 2.15 mL + H2O 2 mL"],
            "products": [
                {"to_slots": ["modulator_1", "modulator_2"],
                 "name": "triethylamine", "amount_text": "2.15 mL"},
                {"to_slots": ["solvent_main", "solvent_secondary"],
                 "name": "water", "amount_text": "2 mL"},
            ],
        }]}

    def test_explicit_composite_requires_complete_expression_and_preserves_origin(self):
        rules = self.composite_rule()
        p.validate_record_rules({"P001-R001": rules}, {"P001-R001"})
        expected = reference_row(modulator_1="triethylamine", modulator_1_amount_text="2.15 mL",
                                 solvent_main="DMF", solvent_main_amount_text="5 mL",
                                 solvent_secondary="water", solvent_secondary_amount_text="2 mL")
        candidate = reference_row(modulator_2="triethylamine + water",
                                  modulator_2_amount_text="TEA 2.15 mL + H2O 2 mL",
                                  solvent_secondary="DMF", solvent_secondary_amount_text="5 mL")
        aligned, origins = p.align_role_slots(candidate, expected, record_rules=rules)
        self.assertEqual(aligned["solvent_secondary"], "water")
        self.assertEqual(aligned["solvent_secondary_amount_text"], "2 mL")
        self.assertEqual(aligned["modulator_1"], "triethylamine")
        self.assertEqual(origins["modulator_1"], "modulator_2")
        self.assertEqual(origins["solvent_secondary"], "modulator_2")
        self.assertEqual(set(aligned["_composite_projection_slots"]),
                         {"modulator_1", "solvent_secondary"})
        self.assertEqual(p.pair_score(candidate, expected, record_rules=rules),
                         p.pair_score(expected, expected))
        for name, amount in (("triethylamine + water", "TEA 2.15 mL"),
                             ("triethylamine + water", "TEA 2.15 mL + H2O 3 mL"),
                             ("triethylamine", "TEA 2.15 mL + H2O 2 mL")):
            changed = dict(candidate, modulator_2=name, modulator_2_amount_text=amount)
            unprojected, _ = p.apply_role_reassignments(changed, rules)
            self.assertEqual(unprojected["modulator_2"], name)
            self.assertEqual(unprojected["solvent_main"], "")
            self.assertFalse(unprojected["_composite_projection_slots"])

    def test_composite_is_atomic_and_does_not_duplicate_an_existing_product(self):
        candidate = reference_row(modulator_1="triethylamine + water",
                                  modulator_1_amount_text="TEA 2.15 mL + H2O 2 mL",
                                  solvent_main="DMF", solvent_main_amount_text="5 mL",
                                  solvent_secondary="water", solvent_secondary_amount_text="2 mL")
        for occupied_name in ("water", "ethanol"):
            candidate["solvent_secondary"] = occupied_name
            unchanged, _ = p.apply_role_reassignments(candidate, self.composite_rule())
            self.assertEqual(unchanged["modulator_1"], "triethylamine + water")
            self.assertEqual(unchanged["modulator_1_amount_text"], "TEA 2.15 mL + H2O 2 mL")
            self.assertEqual(unchanged["solvent_secondary"], occupied_name)
            self.assertFalse(unchanged["_composite_projection_slots"])

    def test_cross_role_linker_transfer_consumes_the_complete_pair_once(self):
        rules = {"role_reassignments": [{
            "names": ["1,7-phenanthroline"], "from_slots": ["modulator_1", "modulator_2"],
            "to_slots": ["linker_1", "linker_2", "linker_3"],
        }]}
        candidate = reference_row(linker_3="terephthalic acid", linker_3_amount_text="10 mg",
                                  modulator_2="1,7-phenanthroline", modulator_2_amount_text="20 mg")
        expected = reference_row(linker_1="terephthalic acid", linker_1_amount_text="10 mg",
                                 linker_2="1,7-phenanthroline", linker_2_amount_text="20 mg")
        aligned, origins = p.align_role_slots(candidate, expected, record_rules=rules)
        self.assertEqual(aligned["linker_2"], "1,7-phenanthroline")
        self.assertEqual(aligned["linker_2_amount_text"], "20 mg")
        self.assertEqual(origins["linker_2"], "modulator_2")
        self.assertFalse(aligned["modulator_1"] or aligned["modulator_2"])

    def test_permitted_role_transfer_preserves_pair_and_source_slot(self):
        expected = reference_row(solvent_main="DMF", solvent_main_amount_text="1 mL",
                                 solvent_secondary="water", solvent_secondary_amount_text="54 uL")
        candidate = reference_row(solvent_main="DMF", solvent_main_amount_text="1 mL",
                                  modulator_1="H2O", modulator_1_amount_text="0.054 mL")
        rules = {"role_reassignments": [{
            "names": ["water"], "from_slots": ["modulator_1", "modulator_2"],
            "to_slots": ["solvent_main", "solvent_secondary"],
        }]}
        aligned, mapping = p.align_role_slots(candidate, expected, record_rules=rules)
        self.assertEqual(mapping["solvent_secondary"], "modulator_1")
        self.assertEqual(aligned["solvent_secondary_amount_text"], "0.054 mL")
        self.assertEqual(aligned["modulator_1"], "")
        self.assertEqual(p.pair_score(candidate, expected, record_rules=rules),
                         p.pair_score(expected, expected, record_rules=rules))
        candidate["solvent_secondary"] = "ethanol"
        candidate["solvent_secondary_amount_text"] = "3 mL"
        unchanged, _ = p.apply_role_reassignments(candidate, rules)
        self.assertEqual(unchanged["modulator_1"], "H2O")
        self.assertEqual(unchanged["solvent_secondary"], "ethanol")

    def test_duplicate_identity_constraint_applies_only_to_explicit_records(self):
        expected = reference_row(solvent_main="DMF ||| ethanol ||| water",
                                 solvent_secondary="DMF ||| ethanol ||| water")
        candidate = reference_row(solvent_main="water", solvent_secondary="H2O")
        self.assertEqual(p.rule_field_matches(candidate, expected), {})
        decisions = p.rule_field_matches(candidate, expected, record_rules={
            "distinct_components": [["solvent_main", "solvent_secondary"]]})
        self.assertEqual(decisions, {"solvent_secondary": (False, "duplicate_component")})
        label = p.field_label("H2O", expected["solvent_secondary"], "solvent_secondary",
                              override_match=decisions["solvent_secondary"][0])
        self.assertEqual((label["TP"], label["FP"], label["FN"]), (0, 1, 1))

    def test_explicit_volume_pair_can_swap_but_wrong_total_or_name_cannot(self):
        expected = reference_row(solvent_main="water", solvent_secondary="DMF")
        rules = {"amount_pair_options": [{
            "slots": ["solvent_main", "solvent_secondary"],
            "values_ml": [[2/3, 4/3]], "absolute_tolerance_ml": 0.01,
        }]}
        candidate = reference_row(solvent_main="water", solvent_main_amount_text="0.67 mL",
                                  solvent_secondary="DMF", solvent_secondary_amount_text="1330 uL")
        self.assertEqual(len(p.rule_field_matches(candidate, expected, record_rules=rules)), 2)
        candidate["solvent_main_amount_text"], candidate["solvent_secondary_amount_text"] = (
            candidate["solvent_secondary_amount_text"], candidate["solvent_main_amount_text"])
        self.assertEqual(len(p.rule_field_matches(candidate, expected, record_rules=rules)), 2)
        candidate["solvent_main_amount_text"] = "8 mL"
        self.assertEqual(p.rule_field_matches(candidate, expected, record_rules=rules), {})
        candidate["solvent_main_amount_text"] = "1.33 mL"
        candidate["solvent_secondary"] = "chloroform"
        self.assertEqual(p.rule_field_matches(candidate, expected, record_rules=rules), {})

    def test_invalid_rule_ids_or_reassignment_slots_are_rejected(self):
        with self.assertRaises(ValueError):
            p.validate_record_rules({"P999-R001": {}}, {"P001-R001"})
        with self.assertRaises(ValueError):
            p.validate_record_rules({"P001-R001": {"role_reassignments": [{
                "names": ["water"], "from_slots": ["solvent_main"],
                "to_slots": ["solvent_main"]}]}}, {"P001-R001"})
        rules = self.composite_rule()
        rules["composite_projections"][0]["products"][1]["amount_text"] = ""
        with self.assertRaises(ValueError):
            p.validate_record_rules({"P001-R001": rules}, {"P001-R001"})


class IntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.reference = self.root / "ground_truth.xlsx"
        self.csv = self.root / "predictions.csv"

    def write_reference(self, rows, papers=None, annotator_rows=None):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet in p.SHEETS:
            tab = workbook.create_sheet(sheet)
            tab.append(list(p.IDENTIFIERS + p.FIELDS))
            for row in (annotator_rows or {}).get(sheet, rows):
                tab.append([row[field] for field in p.IDENTIFIERS + p.FIELDS])
        workbook.save(self.reference)
        if papers is not None:
            p.save_csv(self.root / "papers.csv", papers)

    def write_predictions(self, rows):
        p.save_csv(self.csv, rows)

    def test_validation_cli_uses_configured_defaults_without_creating_outputs(self):
        self.write_reference([reference_row(metal_1="ZnCl2")])
        self.write_predictions([reference_row(metal_1="ZnCl2")])
        output_root = self.root / "results"
        printed = io.StringIO()
        with patch.multiple(p, DEFAULT_CSV=self.csv, DEFAULT_REFERENCE=self.reference,
                            DEFAULT_OUTPUT_ROOT=output_root), redirect_stdout(printed):
            self.assertEqual(p.main(["--validate-only"]), 0)
        result = json.loads(printed.getvalue())
        self.assertEqual(result["reference_records"], 1)
        self.assertEqual(result["sampled_papers"], 1)
        self.assertEqual(result["usable_prediction_records"], 1)
        self.assertFalse(output_root.exists())

    def test_cli_explicit_inputs_allocate_numbered_runs_without_overwriting(self):
        self.write_reference([reference_row(metal_1="ZnCl2", time_h="12")])
        self.write_predictions([reference_row(metal_1="ZnCl2", time_h="12")])
        output_root = self.root / "results"
        args = [str(self.csv), "--ground-truth", str(self.reference),
                "--output-root", str(output_root), "--bootstraps", "0"]
        with redirect_stdout(io.StringIO()):
            self.assertEqual(p.main(args), 0)
            previous = {path.name: path.read_bytes()
                        for path in (output_root / "run_001").iterdir()}
            self.assertEqual(p.main(args), 0)
        self.assertEqual({path.name for path in output_root.iterdir()}, {"run_001", "run_002"})
        self.assertEqual(previous, {path.name: path.read_bytes()
                                   for path in (output_root / "run_001").iterdir()})
        result = json.loads((output_root / "run_002/summary.json").read_text(encoding="utf-8"))
        self.assertEqual(result["record_detection"]["exact_records_26_fields"], 1)
        self.assertEqual(set(result["field_evaluation"]), {"all_records", "matched_records"})
        self.assertEqual(result["scope_labels"], {
            "all_records": "all-protocol evaluation",
            "matched_records": "recovered-protocol evaluation",
        })
        self.assertNotIn("end_to_end", json.dumps(result))
        metrics = p.read_csv(output_root / "run_002/field_metrics.csv")[1]
        self.assertEqual(len(metrics), 28)
        self.assertEqual({row["field"] for row in metrics}, set(p.FIELD_GROUPS))
        self.assertEqual({row["scope"] for row in metrics}, {"all_records", "matched_records"})
        labels = p.read_csv(output_root / "run_002/field_labels.csv")[1]
        self.assertEqual({row["field"] for row in labels}, set(p.FIELDS))
        self.assertEqual(len(labels), 26)
        manifest = json.loads((output_root / "run_002/run_manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["field_groups"], result["field_groups"])
        self.assertEqual(manifest["scope_labels"], result["scope_labels"])

    def test_sparse_annotator_sheets_do_not_change_inventory_or_scores(self):
        first = reference_row(metal_1="ZnCl2", time_h="12")
        second = reference_row("P002-R001", "10.example/b", "MOF-B", time_h="24")
        self.write_predictions([first, second])
        self.write_reference([first, second])
        before_folder, before = p.evaluate(
            self.csv, self.reference, output_dir=self.root / "complete", bootstraps=0)
        self.write_reference([second, first], annotator_rows={
            "A1": [first], "A2": [second, first], "A3": [], "A4": [second],
        })
        after_folder, after = p.evaluate(
            self.csv, self.reference, output_dir=self.root / "sparse", bootstraps=0)
        for section in ("record_detection", "field_evaluation", "paper_evaluation"):
            self.assertEqual(before[section], after[section])
        for filename, keys in (
                ("record_labels.csv", ("record_id", "prediction_id")),
                ("field_labels.csv", ("record_id", "prediction_id", "field"))):
            before_rows = p.read_csv(before_folder / filename)[1]
            after_rows = p.read_csv(after_folder / filename)[1]
            key = lambda row: tuple(row[item] for item in keys)
            self.assertEqual(sorted(before_rows, key=key), sorted(after_rows, key=key))

    def test_new_record_with_one_contributor_enters_reference_denominators(self):
        existing = reference_row(metal_1="ZnCl2", time_h="12")
        added = reference_row("P001-R002", name="MOF-B", metal_1="CuCl2", time_h="24")
        self.write_reference([existing, added], annotator_rows={
            "A1": [existing], "A2": [added], "A3": [], "A4": [],
        })
        self.write_predictions([existing])
        folder, result = p.evaluate(
            self.csv, self.reference, output_dir=self.root / "added", bootstraps=0)
        records = result["record_detection"]
        self.assertEqual((records["reference_records"], records["TP"], records["FN"]), (2, 1, 1))
        labels = p.read_csv(folder / "record_labels.csv")[1]
        self.assertEqual(next(row["record_label"] for row in labels
                              if row["record_id"] == "P001-R002"), "FN")
        self.assertEqual(result["paper_evaluation"]["evaluated_papers"], 1)

    def test_sparse_annotator_sheets_reject_unknown_duplicate_and_wrong_paper_ids(self):
        row = reference_row(metal_1="ZnCl2")
        cases = (
            ([reference_row("P001-R999")], "absent from the primary sheet"),
            ([row, row], "Duplicate record IDs"),
            ([dict(row, paper_id="P002")], "paper identifiers differ"),
            ([dict(row, doi="10.example/other")], "paper identifiers differ"),
        )
        for annotator_rows, error in cases:
            with self.subTest(error=error, rows=annotator_rows):
                self.write_reference([row], annotator_rows={"A2": annotator_rows})
                with self.assertRaisesRegex(ValueError, error):
                    p.load_reference(self.reference)

    def test_missing_paper_extra_record_and_wrong_field_enter_all_records(self):
        self.write_reference([
            reference_row(time_h="12", metal_1="ZnCl2"),
            reference_row("P002-R001", "10.example/b", "MOF-B", time_h="24"),
        ])
        self.write_predictions([
            {"doi": "10.example/a", "mof_name": "MOF-A", "time_h": "99", "metal_1": "ZnCl2"},
            {"doi": "10.example/a", "mof_name": "completely unrelated", "time_h": "400"},
        ])
        folder, result = p.evaluate(self.csv, self.reference, output_dir=self.root / "run", bootstraps=20)
        record = result["record_detection"]
        self.assertEqual((record["TP"], record["FP"], record["FN"]), (1, 1, 1))
        _, metrics = p.read_csv(folder / "field_metrics.csv")
        time = next(row for row in metrics if row["field"] == "time_h" and row["scope"] == "all_records")
        self.assertEqual((int(time["TP"]), int(time["FP"]), int(time["FN"])), (0, 2, 2))
        self.assertEqual(result["paper_evaluation"]["macro_paper_recall"], 0.5)
        manifest = json.loads((folder / "run_manifest.json").read_text())
        self.assertEqual(manifest["input_csv"]["sha256"], p.digest(self.csv))
        self.assertTrue(manifest["completed"])
        with self.assertRaises(FileExistsError):
            p.evaluate(self.csv, self.reference, output_dir=folder, bootstraps=0)

    def test_single_rules_bundle_and_readable_optional_reference(self):
        row = reference_row(solvent_main="DMF", solvent_main_amount_text="5 mL",
                            solvent_secondary="not reported | water")
        self.write_reference([row])
        bundle = {
            "schema_version": "positive-reference-v2",
            "papers": [{"paper_id": "P001", "doi": "10.example/a",
                        "included": True, "exclusion_reason": ""}],
            "accepted_variants": {"P001-R001": {"solvent_main": ["dimethylformamide"]}},
            "role_pair_variants": {}, "record_rules": {}, "rejected_candidates": [],
        }
        (self.root / "reference_rules.json").write_text(json.dumps(bundle), encoding="utf-8")
        self.write_predictions([{"doi": "10.example/a", "mof_name": "MOF-A",
                                 "solvent_main": "dimethylformamide",
                                 "solvent_main_amount_text": "5 mL"}])
        folder, result = p.evaluate(self.csv, self.reference, output_dir=self.root / "bundle",
                                    bootstraps=0)
        self.assertEqual(result["record_detection"]["exact_records_26_fields"], 1)
        self.assertEqual(result["field_evaluation"]["matched_records"]["micro"]["optional_reference_cells"], 1)
        manifest = json.loads((folder / "run_manifest.json").read_text())
        self.assertEqual(set(manifest["reference_companions"]), {"reference_rules"})
        used = json.loads((folder / "reference_rules_used.json").read_text())
        self.assertEqual(used["papers"], bundle["papers"])

    def test_entire_metrics_and_prediction_ids_are_role_order_invariant_with_extras(self):
        row = reference_row(time_h="24", temperature_c="120")
        for group in p.ROLE_GROUPS:
            for position, slot in enumerate(group, 1):
                row[slot] = slot + "-identity"
                row[slot + "_amount_text"] = f"{position} mmol"
        self.write_reference([row])
        extra = reference_row(name="unrelated", metal_3="extra-metal",
                              linker_2="extra-linker", solvent_secondary="extra-solvent")
        predictions = [row, extra]
        self.write_predictions(predictions)
        first, before = p.evaluate(self.csv, self.reference, output_dir=self.root / "before",
                                    bootstraps=0)
        reordered = []
        for candidate in predictions:
            new = dict(candidate)
            for group in p.ROLE_GROUPS:
                for destination, source in zip(group, group[1:] + group[:1]):
                    new[destination] = candidate[source]
                    new[destination + "_amount_text"] = candidate[source + "_amount_text"]
            reordered.append(new)
        self.write_predictions(reordered)
        second, after = p.evaluate(self.csv, self.reference, output_dir=self.root / "after",
                                   bootstraps=0)
        self.assertEqual(before["record_detection"]["TP"], 1)
        self.assertEqual(before["record_detection"]["FP"], 1)
        for section in ("record_detection", "field_evaluation", "paper_evaluation"):
            self.assertEqual(before[section], after[section])
        self.assertEqual(p.read_csv(first / "field_metrics.csv")[1],
                         p.read_csv(second / "field_metrics.csv")[1])
        self.assertEqual({r["prediction_id"] for r in p.read_csv(first / "predictions_used.csv")[1]},
                         {r["prediction_id"] for r in p.read_csv(second / "predictions_used.csv")[1]})

    def test_duplicate_predictions_cannot_recover_two_records(self):
        self.write_reference([reference_row(time_h="12")])
        row = {"doi": "10.example/a", "mof_name": "MOF-A", "time_h": "12"}
        self.write_predictions([row, row])
        _, result = p.evaluate(self.csv, self.reference, output_dir=self.root / "run", bootstraps=0)
        self.assertEqual(result["record_detection"]["TP"], 1)
        self.assertEqual(result["record_detection"]["FP"], 1)

    def test_order_and_full_corpus_selection_do_not_change_metrics(self):
        self.write_reference([reference_row(time_h="12"), reference_row(
            "P002-R001", "10.example/b", "MOF-B", time_h="24")])
        rows = [
            {"doi": "10.example/a", "mof_name": "MOF-A", "time_h": "12"},
            {"doi": "10.example/b", "mof_name": "MOF-B", "time_h": "24"},
        ]
        self.write_predictions(rows)
        first, result1 = p.evaluate(self.csv, self.reference, output_dir=self.root / "first", bootstraps=0)
        self.write_predictions(list(reversed(rows)) + [
            {"doi": "10.example/outside", "mof_name": "Other", "time_h": "36"}])
        second, result2 = p.evaluate(self.csv, self.reference, output_dir=self.root / "second", bootstraps=0)
        self.assertEqual(result1["record_detection"], result2["record_detection"])
        self.assertEqual(result1["field_evaluation"], result2["field_evaluation"])
        ids1 = {row["prediction_id"] for row in p.read_csv(first / "record_labels.csv")[1]}
        ids2 = {row["prediction_id"] for row in p.read_csv(second / "record_labels.csv")[1]}
        self.assertEqual(ids1, ids2)

    def test_failed_rows_and_empty_placeholders_are_not_false_positive_syntheses(self):
        self.write_reference([reference_row(time_h="12")])
        self.write_predictions([
            {"doi": "10.example/a", "mof_name": "", "time_h": "", "status": "ok"},
            {"doi": "10.example/a", "mof_name": "MOF-A", "time_h": "12", "status": "failed"},
        ])
        _, result = p.evaluate(self.csv, self.reference, output_dir=self.root / "run", bootstraps=0)
        self.assertEqual(result["record_detection"]["FN"], 1)
        self.assertEqual(result["record_detection"]["FP"], 0)
        self.assertEqual(result["coverage"]["ignored_rows_by_reason"],
                         {"empty_synthesis_placeholder": 1, "failed_or_invalid_status": 1})

    def test_explicit_exclusion_removes_paper_from_all_denominators(self):
        self.write_reference([
            reference_row(time_h="12"),
            reference_row("P002-R001", "10.example/b", "MOF-B", time_h="24"),
        ], papers=[
            {"paper_id": "P001", "doi": "10.example/a", "included": True, "exclusion_reason": ""},
            {"paper_id": "P002", "doi": "10.example/b", "included": False, "exclusion_reason": "unavailable"},
        ])
        self.write_predictions([{"doi": "10.example/a", "mof_name": "MOF-A", "time_h": "12"}])
        _, result = p.evaluate(self.csv, self.reference, output_dir=self.root / "run", bootstraps=0)
        self.assertEqual(result["record_detection"]["FN"], 0)
        self.assertEqual(result["paper_evaluation"]["evaluated_papers"], 1)
        self.assertEqual(result["paper_evaluation"]["excluded_papers"], 1)

    def test_record_rules_are_loaded_used_and_snapshotted(self):
        self.write_reference([reference_row(
            solvent_main="DMF", solvent_main_amount_text="1 mL",
            solvent_secondary="water", solvent_secondary_amount_text="54 uL")])
        rules = {"P001-R001": {"role_reassignments": [{
            "names": ["water"], "from_slots": ["modulator_1", "modulator_2"],
            "to_slots": ["solvent_main", "solvent_secondary"]}]}}
        (self.root / "record_rules.json").write_text(json.dumps(rules), encoding="utf-8")
        self.write_predictions([{"doi": "10.example/a", "mof_name": "MOF-A",
                                 "solvent_main": "DMF", "solvent_main_amount_text": "1 mL",
                                 "modulator_1": "H2O", "modulator_1_amount_text": "0.054 mL"}])
        folder, result = p.evaluate(self.csv, self.reference, output_dir=self.root / "rules",
                                    bootstraps=0)
        self.assertEqual(result["record_detection"]["exact_records_26_fields"], 1)
        self.assertEqual(json.loads((folder / "reference_rules_used.json").read_text())["record_rules"], rules)
        labels = p.read_csv(folder / "field_labels.csv")[1]
        water = next(row for row in labels if row["field"] == "solvent_secondary")
        self.assertEqual(water["original_prediction_slot"], "modulator_1")
        self.assertEqual(water["record_rule"], "permitted_role_reassignment")

    def test_composite_projection_labels_original_expression_and_is_order_invariant(self):
        expected = reference_row(
            modulator_1="triethylamine", modulator_1_amount_text="2.15 mL",
            solvent_main="DMF", solvent_main_amount_text="5 mL",
            solvent_secondary="water", solvent_secondary_amount_text="2 mL")
        self.write_reference([expected])
        rules = {"P001-R001": RecordRuleTests.composite_rule()}
        (self.root / "record_rules.json").write_text(json.dumps(rules), encoding="utf-8")
        candidate = reference_row(
            modulator_1="triethylamine + water",
            modulator_1_amount_text="TEA 2.15 mL + H2O 2 mL",
            solvent_main="DMF", solvent_main_amount_text="5 mL")
        self.write_predictions([candidate])
        folder, before = p.evaluate(self.csv, self.reference,
                                     output_dir=self.root / "composite", bootstraps=0)
        self.assertEqual(before["record_detection"]["exact_records_26_fields"], 1)
        labels = p.read_csv(folder / "field_labels.csv")[1]
        water = next(row for row in labels if row["field"] == "solvent_secondary_amount_text")
        self.assertEqual(water["prediction_value"], "2 mL")
        self.assertEqual(water["original_prediction_value"], "TEA 2.15 mL + H2O 2 mL")
        self.assertEqual(water["original_prediction_slot"], "modulator_1")
        self.assertEqual(water["record_rule"], "explicit_composite_projection")
        rotated = dict(candidate)
        for group in p.ROLE_GROUPS:
            for destination, source in zip(group, group[1:] + group[:1]):
                for suffix in ("", "_amount_text"):
                    rotated[destination + suffix] = candidate[source + suffix]
        self.write_predictions([rotated])
        _, after = p.evaluate(self.csv, self.reference,
                               output_dir=self.root / "composite_rotated", bootstraps=0)
        self.assertEqual(before["record_detection"], after["record_detection"])
        self.assertEqual(before["field_evaluation"], after["field_evaluation"])

    def test_cross_role_rule_does_not_discard_duplicate_water(self):
        self.write_reference([reference_row(solvent_main="DMF", solvent_secondary="water")])
        rules = {"P001-R001": {"role_reassignments": [{
            "names": ["water"], "from_slots": ["modulator_1", "modulator_2"],
            "to_slots": ["solvent_main", "solvent_secondary"]}]}}
        (self.root / "record_rules.json").write_text(json.dumps(rules), encoding="utf-8")
        self.write_predictions([{"doi": "10.example/a", "mof_name": "MOF-A",
                                 "solvent_main": "DMF", "solvent_secondary": "water",
                                 "modulator_1": "water"}])
        folder, _ = p.evaluate(self.csv, self.reference, output_dir=self.root / "duplicate",
                               bootstraps=0)
        labels = p.read_csv(folder / "field_labels.csv")[1]
        extra = next(row for row in labels if row["field"] == "modulator_1")
        self.assertEqual(extra["label"], "extra_value")
        self.assertEqual(extra["FP"], "1")


if __name__ == "__main__":
    unittest.main()
