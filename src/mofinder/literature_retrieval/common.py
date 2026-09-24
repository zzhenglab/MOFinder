"""Portable literature retrieval settings, neutral profiles, and local inventories."""

from collections import Counter
import hashlib
import json
import math
from pathlib import Path
import re
import shutil


PUBLISHER_IDS = ("publisher_W", "publisher_A", "publisher_R", "publisher_S", "publisher_E")
ICON_EXTENSIONS = (".png", ".jpg", ".jpeg", ".bmp")
SI_TEMPLATES = {
    "publisher_W": ["1", "2", "3a", "3aa", "3b", "3bb", "4"],
    "publisher_A": ["1", "1b", "1c", "2a", "2aa", "2b", "2c", "2d", "3"],
    "publisher_R": ["1", "2"],
    "publisher_S": ["Accept", "Accept2", "1", "2", "2b", "2c", "3b", "3"],
    "publisher_E": ["Accept", "Accept2", "Accept3", "1", "1b", "2", "2b", "3"],
}


def publisher_key(value):
    """Resolve an explicit neutral profile without inferring it from a DOI."""
    if not isinstance(value, str):
        return None
    return next((key for key in PUBLISHER_IDS if value.strip().casefold() == key.casefold()), None)


def load_settings(config_file, mode):
    """Resolve configuration paths without creating files or opening a desktop."""
    if mode not in {"papers", "si"}:
        raise ValueError("Literature retrieval mode must be papers or si.")
    path = Path(config_file).expanduser().resolve()
    config = json.loads(path.read_text(encoding="utf-8"))
    root = (path.parent / config.get("project_root", "..")).resolve()
    section = config[mode]
    paths = {key: (root / Path(section[key]).expanduser()).resolve() for key in
             ["workbook", "calibration_file", "app_settings_file", "download_dir", "working_dir"]}
    paths["paper_processing_icon_dir"] = (root / Path(config["paper_processing_icon_dir"]).expanduser()).resolve()
    settings = {"mode": mode, "config_file": path, "project_root": root, **paths,
                "sample_urls": section.get("sample_urls", {}), "tuning": section.get("tuning", {})}
    for key, value in settings["sample_urls"].items():
        if publisher_key(key) is None or not isinstance(value, str) or not value.startswith("https://"):
            raise ValueError("Sample URLs must use neutral profile IDs and HTTPS URLs.")
    for key, value in settings["tuning"].items():
        if key == "APPLY_HYPERLINKS_ON_FINAL_SAVE":
            if not isinstance(value, bool):
                raise ValueError(f"{key} must be true or false.")
            continue
        if type(value) not in (int, float) or not math.isfinite(value) or value < 0:
            raise ValueError(f"Invalid nonnegative timing or limit setting: {key}")
        if key in {"SAVE_EVERY", "MAX_RETRIES", "MAX_JOURNAL_FAILS", "MAX_RETRIES_PER_ROW",
                   "SAVE_EVERY_PENDING", "MAX_JOURNAL_FAILS_BASE", "JOURNAL_SUCCESS_BOOST_MULTIPLIER",
                   "PUBLISHER_W_SCROLL_UP_COUNT"} and (type(value) is not int or value < 1):
            raise ValueError(f"{key} must be a positive integer.")
        if key in {"ICON_SEARCH_DELAY", "ICON_SEARCH_TICK", "PUBLISHER_W_SCROLL_STEP_SEC"} and value <= 0:
            raise ValueError(f"{key} must be greater than zero.")
        if key == "ICON_CONFIDENCE" and not 0 < value <= 1:
            raise ValueError("ICON_CONFIDENCE must be greater than zero and at most one.")
    return settings


def read_inventory(path):
    """Read a CSV/XLSX inventory while retaining status strings and blank cells."""
    import pandas as pd
    path = Path(path)
    if path.suffix.lower() == ".csv":
        frame = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    elif path.suffix.lower() == ".xlsx":
        frame = pd.read_excel(path, dtype=str, keep_default_na=False)
    else:
        raise ValueError("Use a .csv or .xlsx literature retrieval inventory.")
    if "DOI" not in frame.columns:
        raise ValueError("Missing column: DOI")
    # An explicit neutral ID column takes precedence over descriptive metadata.
    if "Publisher ID" in frame.columns:
        frame["Publisher"] = frame["Publisher ID"]
    if "Publisher" not in frame.columns:
        raise ValueError("Missing column: Publisher (neutral ID), or Publisher ID")
    return frame


def write_inventory(frame, path, *, hyperlinks=True):
    """Atomically save a working inventory; only spreadsheet files need links."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + ".tmp" + path.suffix)
    try:
        if path.suffix.lower() == ".csv":
            frame.to_csv(temporary, index=False, encoding="utf-8-sig")
        elif path.suffix.lower() == ".xlsx":
            # Export the current inventory and optional DOI links.
            frame.to_excel(temporary, index=False)
            if hyperlinks and "DOI Link" in frame.columns:
                from openpyxl import load_workbook
                book = load_workbook(temporary)
                sheet = book.active
                column = list(frame.columns).index("DOI Link") + 1
                for row in range(2, sheet.max_row + 1):
                    cell = sheet.cell(row, column)
                    if isinstance(cell.value, str) and cell.value.strip():
                        cell.hyperlink = cell.value
                        cell.style = "Hyperlink"
                book.save(temporary)
                book.close()
        else:
            raise ValueError("Use a .csv or .xlsx literature retrieval inventory.")
        temporary.replace(path)
    finally:
        if temporary.exists():
            temporary.unlink()


def prepare_local_inventory(settings):
    """Create or reopen a local working copy without modifying archived inputs.

    Source path and hash identify a working copy, so choosing a new source cannot
    silently reuse the status table of another selection. Existing progress is
    retained when the same source is selected again.
    """
    source = Path(settings["workbook"]).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Missing literature retrieval inventory: {source}")
    if source.suffix.lower() not in {".csv", ".xlsx"}:
        raise ValueError("Use a .csv or .xlsx literature retrieval inventory.")
    folder = Path(settings["working_dir"]).resolve()
    folder.mkdir(parents=True, exist_ok=True)
    # Selecting an existing working file continues that exact file.
    if source.is_relative_to(folder):
        return source
    identity = hashlib.sha256(str(source).encode() + b"\0" + source.read_bytes()).hexdigest()[:12]
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", source.stem).strip("_") or "inventory"
    destination = folder / f"{stem}_{identity}{source.suffix.lower()}"
    if not destination.exists():
        shutil.copy2(source, destination)
        provenance = {"source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                      "working_file": str(destination), "mode": settings["mode"]}
        destination.with_suffix(destination.suffix + ".source.json").write_text(
            json.dumps(provenance, indent=2) + "\n", encoding="utf-8")
    return destination


def audit_inventory(path, mode, icon_dir=None):
    """Report records, status coverage, and referenced icons without modifying input."""
    if mode not in {"papers", "si"}:
        raise ValueError("Literature retrieval mode must be papers or si.")
    frame = read_inventory(path)
    field = "Downloaded" if mode == "papers" else "SI Downloaded"
    values = frame[field].tolist() if field in frame else [""] * len(frame)
    if mode == "si":
        def normalize(value):
            value = str(value).strip().lower()
            return "1" if value in {"1", "1.0", "true", "yes", "y"} else (
                   "0" if value in {"0", "0.0", "false", "no", "n"} else "")
        values = [normalize(v) for v in values]
    else:
        values = [str(v).strip() for v in values]
    profiles = [publisher_key(value) for value in frame["Publisher"]]
    ids = frame["DOI"].astype(str).str.strip()
    report = {
        "mode": mode, "inventory": str(Path(path).resolve()), "rows": len(frame),
        "status_column": field, "status_counts": dict(Counter(values)),
        "pending_rows": sum(v == "" for v in values),
        "retry_rows": sum(v == "0" for v in values),
        "publisher_counts": dict(Counter(key or "unmapped" for key in profiles)),
        "unmapped_rows": sum(key is None for key in profiles),
        "unmapped_pending_rows": sum(key is None and status == "" for key, status in zip(profiles, values)),
        "duplicate_doi_rows": int(ids.duplicated().sum()),
        "blank_doi_rows": int(ids.eq("").sum()),
    }
    if icon_dir is not None:
        folder = Path(icon_dir)
        available = {p.stem for p in folder.iterdir() if p.is_file() and p.suffix.lower() in ICON_EXTENSIONS} if folder.is_dir() else set()
        report["icon_directory"] = str(folder.resolve())
        report["icon_templates"] = len(available)
        if mode == "papers":
            report["article_icons_per_profile"] = {
                key: sum(bool(re.fullmatch(re.escape(key) + r"_\d+", stem, re.I)) for stem in available)
                for key in PUBLISHER_IDS}
        else:
            report["missing_referenced_templates"] = sorted(
                key + "_SI_" + suffix for key, suffixes in SI_TEMPLATES.items()
                for suffix in suffixes if key + "_SI_" + suffix not in available)
            report["template_note"] = "Alternative templates are used conditionally; missing templates are reported, not synthesized."
    return report
