"""Evaluate saved positive-synthesis extractions against the reviewed reference.

Run ``python -m mofinder.evaluation.positive [csv]`` to score a raw extraction
CSV. Defaults use the bundled 200-paper extraction and benchmark workbook.
No model requests are made. Each analysis saves record and field labels,
metrics, matching diagnostics, reference snapshots, and input hashes.
"""

from __future__ import annotations

import argparse
import base64
import binascii
from collections import Counter, defaultdict
from datetime import datetime, timezone
import csv
import difflib
import gzip
import hashlib
import itertools
import json
import math
from pathlib import Path
import re
import shutil
import unicodedata
from urllib.parse import unquote
import zlib

import numpy as np
from openpyxl import load_workbook
from openpyxl.packaging.custom import StringProperty
from scipy.optimize import linear_sum_assignment

from mofinder.display import display_path
from mofinder.run_paths import create_run_directory


REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_CSV = REPO_ROOT / "data/processed_data/positive_extraction_200.csv"
DEFAULT_REFERENCE = REPO_ROOT / "benchmarks/positive_extraction/ground_truth.xlsx"
DEFAULT_OUTPUT_ROOT = REPO_ROOT / "results/positive_extraction"


SCHEMA_VERSION = "positive-evaluation-v7-self-contained-reference"
FIELDS = tuple(
    field
    for role, count in (("metal", 3), ("linker", 3), ("modulator", 2))
    for number in range(1, count + 1)
    for field in (f"{role}_{number}", f"{role}_{number}_amount_text")
) + (
    "solvent_main", "solvent_main_amount_text",
    "solvent_secondary", "solvent_secondary_amount_text",
    "temperature_c", "time_h", "tga_decomposition_temp_c",
    "BET_surface_area_m2g", "crystal_code", "topology_code",
)
IDENTIFIERS = ("record_id", "paper_id", "doi", "mof_name")
SHEETS = ("Ground_truth_A1toA4", "A1", "A2", "A3", "A4")
REFERENCE_SHEET_NAMES = {
    "ground_truth_a1toa4", "merged", "ground_truth", "consensus",
}
ALTERNATIVE_SEPARATOR = " | "
UNKNOWN = {"__UNRESOLVED__", "__UNREVIEWED__"}
ABSENT = {"", "none", "null", "nan", "n/a", "__absent__", "not reported", "not specified",
          "no hard value reported", "empty"}
NAME_FIELDS = {
    field for field in FIELDS
    if field.startswith(("metal_", "linker_", "modulator_", "solvent_"))
    and not field.endswith("_amount_text")
}
MATCH_ID_FIELDS = (
    "metal_1", "metal_2", "linker_1", "linker_2", "modulator_1",
    "solvent_main", "solvent_secondary",
)
MATCH_AMOUNT_FIELDS = (
    "metal_1_amount_text", "linker_1_amount_text", "modulator_1_amount_text",
    "solvent_main_amount_text", "temperature_c", "time_h",
)
ROLE_GROUPS = (
    ("metal_1", "metal_2", "metal_3"),
    ("linker_1", "linker_2", "linker_3"),
    ("modulator_1", "modulator_2"),
    ("solvent_main", "solvent_secondary"),
)
FIELD_GROUPS = {
    "metal_name": ("metal_1", "metal_2", "metal_3"),
    "metal_amount": ("metal_1_amount_text", "metal_2_amount_text", "metal_3_amount_text"),
    "linker_name": ("linker_1", "linker_2", "linker_3"),
    "linker_amount": ("linker_1_amount_text", "linker_2_amount_text", "linker_3_amount_text"),
    "modulator_name": ("modulator_1", "modulator_2"),
    "modulator_amount": ("modulator_1_amount_text", "modulator_2_amount_text"),
    "solvent_name": ("solvent_main", "solvent_secondary"),
    "solvent_amount": ("solvent_main_amount_text", "solvent_secondary_amount_text"),
    "temperature_c": ("temperature_c",),
    "time_h": ("time_h",),
    "tga_decomposition_temp_c": ("tga_decomposition_temp_c",),
    "BET_surface_area_m2g": ("BET_surface_area_m2g",),
    "crystal_code": ("crystal_code",),
    "topology_code": ("topology_code",),
}
SCOPE_LABELS = {
    "all_records": "all-protocol evaluation",
    "matched_records": "recovered-protocol evaluation",
}
REFERENCE_METADATA_PREFIX = "MOFinder.ReferenceRules."
REFERENCE_METADATA_SCHEMA = "positive-reference-embedded-v1"
REFERENCE_METADATA_ENCODING = "json+gzip+base64-v1"
REFERENCE_COMPONENTS = {
    "papers": "papers.csv",
    "rejected_candidates": "rejected_candidates.csv",
    "accepted_variants": "accepted_variants.json",
    "role_pair_variants": "role_pair_variants.json",
    "record_rules": "record_rules.json",
}
# These are identities, not a claim that metal counterions, hydration states,
# acid/base states, or concentrations are interchangeable.
ALIASES = {
    "n,n-dimethylformamide": "dmf", "dimethylformamide": "dmf",
    "n,n-diethylformamide": "def", "diethylformamide": "def",
    "dimethyl sulfoxide": "dmso", "dimethylsulfoxide": "dmso",
    "ethanol": "etoh", "methanol": "meoh",
    "isopropanol": "ipa", "isopropyl alcohol": "ipa",
    "h2o": "water", "deionized water": "water", "distilled water": "water",
    "deionised water": "water", "di water": "water",
    "acetic acid": "hoac", "glacial acetic acid": "hoac",
    "hydrochloric acid": "hcl", "nitric acid": "hno3",
    "sulfuric acid": "h2so4", "sodium hydroxide": "naoh",
    "potassium hydroxide": "koh",
    "terephthalic acid": "h2bdc", "1,4-benzenedicarboxylic acid": "h2bdc",
    "trimesic acid": "h3btc", "1,3,5-benzenetricarboxylic acid": "h3btc",
    "4,4'-bipyridine": "bipy", "4,4'-bipy": "bipy",
}
NUMBER = r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?"
QUANTITY = re.compile(
    rf"(?P<value>{NUMBER})\s*(?P<unit>mmol|umol|mol|mg|ug|kg|g|ml|ul|l|"
    r"minutes?|mins?|min|hours?|hrs?|hr|h|days?|d|seconds?|secs?|sec|s)"
    r"(?![a-z])", re.IGNORECASE,
)
UNITS = {
    "mol": ("mol", 1.0), "mmol": ("mol", 1e-3), "umol": ("mol", 1e-6),
    "kg": ("g", 1e3), "g": ("g", 1.0), "mg": ("g", 1e-3), "ug": ("g", 1e-6),
    "l": ("ml", 1e3), "ml": ("ml", 1.0), "ul": ("ml", 1e-3),
}
for _unit in ("h", "hr", "hrs", "hour", "hours"):
    UNITS[_unit] = ("h", 1.0)
for _unit in ("min", "mins", "minute", "minutes"):
    UNITS[_unit] = ("h", 1 / 60)
for _unit in ("s", "sec", "secs", "second", "seconds"):
    UNITS[_unit] = ("h", 1 / 3600)
for _unit in ("d", "day", "days"):
    UNITS[_unit] = ("h", 24.0)


def text(value):
    """Return a readable cell without treating zero as missing."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def doi_key(value):
    """Normalize DOI wrappers, case, percent escapes, and surrounding space."""
    value = unquote(text(value)).strip()
    value = re.sub(r"^(?:https?://)?(?:dx\.)?doi\.org/", "", value, flags=re.I)
    value = re.sub(r"^doi:\s*", "", value, flags=re.I)
    return value.strip().lower()


def clean(value):
    value = unicodedata.normalize("NFKC", text(value)).casefold()
    value = value.replace("µ", "u").replace("μ", "u")
    value = value.translate(str.maketrans({"′": "'", "’": "'", "‘": "'",
                                          "–": "-", "−": "-", "‐": "-"}))
    return re.sub(r"\s+", " ", value).strip()


def is_absent(value):
    return clean(value) in ABSENT


def is_unknown(value):
    return text(value).upper() in UNKNOWN


def split_alternatives(value):
    """Read the display separator and the historical triple-pipe separator."""
    return [text(item) for item in re.split(r"\s*\|\|\|\s*|\s+\|\s+", text(value))]


def display_alternatives(value):
    return ALTERNATIVE_SEPARATOR.join(
        "not reported" if item.upper() == "__ABSENT__" else item
        for item in split_alternatives(value))


def display_reference_data(value):
    """Use one readable separator in serialized reference snapshots."""
    if isinstance(value, dict):
        return {key: display_reference_data(item) for key, item in value.items()}
    if isinstance(value, list):
        return [display_reference_data(item) for item in value]
    return display_alternatives(value) if isinstance(value, str) else value


def alternatives(value, additional=()):
    """Read explicit alternatives; semicolons remain part of a chemical value."""
    if is_unknown(value):
        return []
    values = split_alternatives(value) + [
        item for extra in additional for item in split_alternatives(extra)]
    result = []
    for item in values:
        item = text(item)
        if not is_absent(item) and item not in result:
            result.append(item)
    return result


def accepts_absence(value, additional=()):
    """Absence is accepted only for an empty reference or an explicit variant."""
    if is_unknown(value):
        return False
    variants = split_alternatives(value) + [
        item for extra in additional for item in split_alternatives(extra)]
    return (not alternatives(value, additional) or
            any(text(item) and is_absent(item) for item in variants))


def _number(value):
    try:
        value = float(clean(value).replace(",", ""))
        return value if math.isfinite(value) else None
    except ValueError:
        return None


def _quantities(value, field):
    """Normalize explicit units only, without molar-mass or density inference."""
    value = clean(value)
    numeric = _number(value)
    if numeric is not None:
        if field in {"temperature_c", "tga_decomposition_temp_c"}:
            return {"C": [numeric]}
        if field == "time_h":
            return {"h": [numeric]}
        if field == "BET_surface_area_m2g":
            return {"m2/g": [numeric]}
        return {"number": [numeric]}
    if field in {"temperature_c", "tga_decomposition_temp_c"}:
        match = re.fullmatch(
            rf"({NUMBER})\s*(?:°|deg(?:rees?)?\.?\s*)?([ckf])", value)
        if match:
            amount, unit = float(match[1]), match[2]
            return {"C": [amount if unit == "c" else
                           amount - 273.15 if unit == "k" else (amount - 32) * 5 / 9]}
        return {}
    if field == "BET_surface_area_m2g":
        match = re.fullmatch(rf"({NUMBER})\s*m\s*\^?2\s*/?\s*g(?:\s*\^?-?1)?", value)
        return {"m2/g": [float(match[1])]} if match else {}
    found = list(QUANTITY.finditer(value))
    if not found:
        return {}
    # Do not reduce ranges, ordered thermal stages, concentrations or ratios to
    # one convenient scalar. Those require explicit accepted reference variants.
    residue = QUANTITY.sub("", value)
    if re.search(r"\d|[:/]|(?:\b(?:after|then|to|at|for)\b)", residue):
        return {}
    if residue.strip(" ,;()[]") not in {"", "approximately", "about", "~", "≈"}:
        return {}
    quantities = defaultdict(list)
    for match in found:
        dimension, factor = UNITS[match["unit"].lower()]
        quantities[dimension].append(float(match["value"]) * factor)
    return dict(quantities)


def _quantity_equal(candidate, reference, field, relative_tolerance=1e-6):
    left, right = _quantities(candidate, field), _quantities(reference, field)
    shared = left.keys() & right.keys()
    if not shared:
        return False
    # Two units describing the same quantity, e.g. mg and mmol, are alternate
    # representations. Every dimension reported on both sides must agree.
    return all(
        len(left[dimension]) == len(right[dimension])
        and all(math.isclose(a, b, rel_tol=relative_tolerance, abs_tol=1e-9)
                for a, b in zip(sorted(left[dimension]), sorted(right[dimension])))
        for dimension in shared
    )


def canonical_name(value):
    value = clean(value)
    return ALIASES.get(value, value)


def value_matches(candidate, reference, field, additional=()):
    """Compare one field with documented variants and conservative unit rules."""
    accepted = alternatives(reference, additional)
    if is_absent(candidate):
        return accepts_absence(reference, additional)
    for expected in accepted:
        if clean(candidate) == clean(expected):
            return True
        if field in NAME_FIELDS and canonical_name(candidate) == canonical_name(expected):
            return True
        if field not in NAME_FIELDS and _quantity_equal(candidate, expected, field):
            return True
        # Numeric identifiers may be written by Excel as 123.0.
        if field == "crystal_code":
            cn, gn = _number(candidate), _number(expected)
            if cn is not None and gn is not None and cn == gn:
                return True
    return False


def projected_row(row):
    """Reduce the 83-column extraction schema to the 26 evaluation fields."""
    result = {field: text(row.get(field)) for field in FIELDS}
    result.update({key: text(row.get(key)) for key in IDENTIFIERS})
    result["doi"] = doi_key(row.get("doi") or row.get("DOI"))
    for role, count in (("metal", 3), ("linker", 3), ("modulator", 2)):
        for number in range(1, count + 1):
            prefix = f"{role}_{number}"
            if is_absent(result[prefix]):
                result[prefix] = text(row.get(prefix + "_abbr"))
            if is_absent(result[prefix + "_amount_text"]):
                amount = text(row.get(prefix + "_amount_value"))
                unit = text(row.get(prefix + "_amount_unit"))
                result[prefix + "_amount_text"] = f"{amount} {unit}".strip() if amount else ""
    for prefix in ("solvent_main", "solvent_secondary"):
        if is_absent(result[prefix]):
            result[prefix] = text(row.get(prefix + "_abbr"))
        if is_absent(result[prefix + "_amount_text"]):
            amount = text(row.get(prefix + "_ml"))
            result[prefix + "_amount_text"] = f"{amount} mL" if amount else ""
    for numeric, wording in (("temperature_c", "temperature_c_text"), ("time_h", "time_text")):
        if is_absent(result[numeric]):
            result[numeric] = text(row.get(wording))
    return result


def read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames is None or len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise ValueError(f"Missing or duplicate CSV headers: {path}")
        rows = []
        for row_number, row in enumerate(reader, 2):
            if None in row:
                raise ValueError(f"Too many CSV values at line {row_number}: {path}")
            rows.append(row)
        return reader.fieldnames, rows


def save_csv(path, rows, columns=None):
    rows = list(rows)
    if columns is None:
        columns = list(dict.fromkeys(key for row in rows for key in row))
    with Path(path).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns or ["no_rows"])
        writer.writeheader()
        writer.writerows(rows)


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _included(value):
    value = clean(value)
    if value in {"true", "yes", "1", "included", "include"}:
        return True
    if value in {"false", "no", "0", "excluded", "exclude"}:
        return False
    raise ValueError(f"Paper inclusion must be explicit true/false, received {value!r}")


def _validate_embedded_bundle(bundle):
    """Require every component so damaged metadata cannot change scoring defaults."""
    if not isinstance(bundle, dict) or not set(REFERENCE_COMPONENTS).issubset(bundle):
        raise ValueError("Embedded reference metadata must contain all five matching components.")
    for key in REFERENCE_COMPONENTS:
        expected = list if key in ("papers", "rejected_candidates") else dict
        if not isinstance(bundle[key], expected):
            raise ValueError(f"Invalid embedded reference component: {key}.")


def embed_reference_rules(workbook, bundle):
    """Embed a complete bundle in standard XLSX properties, preserving worksheets.

    The caller saves the workbook. Payload chunks use at most 255 characters per
    custom property; the checksum covers the uncompressed UTF-8 JSON.
    """
    _validate_embedded_bundle(bundle)
    payload = json.dumps(bundle, ensure_ascii=False, separators=(",", ":"),
                         allow_nan=False).encode("utf-8")
    encoded = base64.b64encode(gzip.compress(payload, mtime=0)).decode("ascii")
    chunks = [encoded[start:start + 255] for start in range(0, len(encoded), 255)]
    if len(chunks) > 9999:
        raise ValueError("Reference metadata exceeds the supported workbook property count.")
    properties = workbook.custom_doc_props
    for name in list(properties.names):
        if name.startswith(REFERENCE_METADATA_PREFIX):
            del properties[name]
    metadata = {
        "Schema": REFERENCE_METADATA_SCHEMA, "Encoding": REFERENCE_METADATA_ENCODING,
        "ChunkCount": str(len(chunks)), "SHA256": hashlib.sha256(payload).hexdigest(),
    }
    for name, value in metadata.items():
        properties.append(StringProperty(name=REFERENCE_METADATA_PREFIX + name, value=value))
    for number, chunk in enumerate(chunks, 1):
        properties.append(StringProperty(
            name=REFERENCE_METADATA_PREFIX + f"{number:04d}", value=chunk))
    return metadata


def read_embedded_reference_rules(workbook):
    """Return a verified bundle, or None only when its property namespace is absent."""
    properties = {
        prop.name.removeprefix(REFERENCE_METADATA_PREFIX): prop.value
        for prop in workbook.custom_doc_props
        if prop.name.startswith(REFERENCE_METADATA_PREFIX)
    }
    if not properties:
        return None
    required = {"Schema", "Encoding", "ChunkCount", "SHA256"}
    if not required.issubset(properties):
        raise ValueError("Embedded reference metadata is incomplete; restore the complete workbook.")
    if properties["Schema"] != REFERENCE_METADATA_SCHEMA:
        raise ValueError("Unsupported embedded reference metadata schema.")
    if properties["Encoding"] != REFERENCE_METADATA_ENCODING:
        raise ValueError("Unsupported embedded reference metadata encoding.")
    count_text = properties["ChunkCount"]
    if not isinstance(count_text, str) or not re.fullmatch(r"[1-9][0-9]{0,3}", count_text):
        raise ValueError("Invalid embedded reference metadata chunk count.")
    chunk_names = [f"{number:04d}" for number in range(1, int(count_text) + 1)]
    if set(properties) != required | set(chunk_names):
        raise ValueError("Embedded reference metadata chunks are missing or unexpected.")
    chunks = [properties[name] for name in chunk_names]
    if any(not isinstance(chunk, str) or not 1 <= len(chunk) <= 255 for chunk in chunks):
        raise ValueError("Invalid embedded reference metadata chunk.")
    try:
        payload = gzip.decompress(base64.b64decode("".join(chunks), validate=True))
    except (ValueError, binascii.Error, OSError, EOFError, zlib.error) as exc:
        raise ValueError("Embedded reference metadata is corrupt or truncated.") from exc
    if hashlib.sha256(payload).hexdigest() != properties["SHA256"]:
        raise ValueError("Embedded reference metadata checksum mismatch.")
    try:
        bundle = json.loads(payload.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("Embedded reference metadata is not valid UTF-8 JSON.") from exc
    _validate_embedded_bundle(bundle)
    return bundle


def load_reference(path):
    path = Path(path).expanduser().resolve()
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet_names = tuple(workbook.sheetnames)
        if (len(sheet_names) != 5 or sheet_names[1:] != SHEETS[1:]
                or sheet_names[0].casefold() not in REFERENCE_SHEET_NAMES):
            raise ValueError(f"Reference workbook must have exactly these sheets in order: {SHEETS}; "
                             "Merged and Consensus first-sheet names are also supported.")
        records = []
        primary_by_id = {}
        for sheet_name in sheet_names:
            values = workbook[sheet_name].iter_rows(values_only=True)
            headers = tuple(text(cell) for cell in next(values))
            if headers != IDENTIFIERS + FIELDS:
                raise ValueError(f"{sheet_name}: expected identifier columns followed by the 26 fields.")
            ids = []
            for cells in values:
                if not any(cell is not None for cell in cells):
                    continue
                if any(isinstance(cell, str) and cell.startswith("=") for cell in cells):
                    raise ValueError(f"Reference formulas are not supported in {sheet_name}.")
                row = {key: text(cell) for key, cell in zip(headers, cells)}
                for field in FIELDS:
                    row[field] = display_alternatives(row[field])
                row["doi"] = doi_key(row["doi"])
                if not row["record_id"] or not row["paper_id"] or not row["doi"].startswith("10."):
                    raise ValueError(f"Invalid reference identifiers in {sheet_name}: {row['record_id']!r}")
                ids.append(row["record_id"])
                if sheet_name == sheet_names[0]:
                    records.append(row)
                    primary_by_id[row["record_id"]] = row
                else:
                    primary = primary_by_id.get(row["record_id"])
                    if primary is None:
                        raise ValueError(
                            f"{sheet_name}: record ID is absent from the primary sheet: "
                            f"{row['record_id']}")
                    if any(row[key] != primary[key] for key in ("paper_id", "doi")):
                        raise ValueError(
                            f"{sheet_name}: paper identifiers differ from the primary sheet: "
                            f"{row['record_id']}")
            if len(ids) != len(set(ids)):
                raise ValueError(f"Duplicate record IDs in {sheet_name}.")
        if not records:
            raise ValueError("The consensus is empty.")
        bundle = read_embedded_reference_rules(workbook)
    finally:
        workbook.close()
    companion_files = {}
    folder = path.parent
    bundle_path = folder / "reference_rules.json"
    embedded = bundle is not None
    if bundle is None and bundle_path.is_file():
        bundle = json.loads(bundle_path.read_text(encoding="utf-8-sig"))
        if not isinstance(bundle, dict):
            raise ValueError("reference_rules.json must contain an object.")
        companion_files["reference_rules"] = bundle_path
    if bundle is None and not any((folder / name).is_file()
                                  for name in REFERENCE_COMPONENTS.values()):
        raise ValueError(
            "Reference matching metadata is missing. Use the complete benchmark workbook "
            "or an explicit legacy reference companion. For a cell-only reference, embed its "
            "paper manifest with empty matching-rule components using embed_reference_rules().")

    def component(key, filename, default):
        """Prefer embedded metadata; legacy companion files remain readable."""
        if bundle is not None:
            return bundle.get(key, default)
        legacy_path = folder / filename
        if not legacy_path.is_file():
            return default
        companion_files[key] = legacy_path
        return (read_csv(legacy_path)[1] if legacy_path.suffix == ".csv" else
                json.loads(legacy_path.read_text(encoding="utf-8-sig")))

    papers = component("papers", "papers.csv", None)
    if papers is not None:
        if not isinstance(papers, list):
            raise ValueError("The paper manifest must be a list.")
        for row in papers:
            row["doi"] = doi_key(row["doi"])
            row["included"] = _included(row["included"])
    else:
        papers = [{"paper_id": row["paper_id"], "doi": row["doi"],
                   "included": True, "exclusion_reason": ""}
                  for row in {r["doi"]: r for r in records}.values()]
    if len(papers) != len({r["doi"] for r in papers}):
        raise ValueError("Duplicate DOI in papers.csv.")
    paper_map = {row["doi"]: row for row in papers}
    for row in records:
        if row["doi"] not in paper_map or row["paper_id"] != paper_map[row["doi"]]["paper_id"]:
            raise ValueError(f"Reference record does not agree with the paper manifest: {row['record_id']}")
    source = component("rejected_candidates", "rejected_candidates.csv", [])
    rejected = []
    if source:
        rejected = [projected_row(row) for row in source]
        if any(not row["record_id"] or row["doi"] not in paper_map for row in rejected):
            raise ValueError("Rejected candidates require stable IDs and benchmark DOIs.")
    variants = component("accepted_variants", "accepted_variants.json", {})
    if variants:
        valid_ids = {row["record_id"] for row in records + rejected}
        if not isinstance(variants, dict):
            raise ValueError("Accepted variants must be a record-ID mapping.")
        for record_id, fields in variants.items():
            if record_id not in valid_ids or not isinstance(fields, dict):
                raise ValueError(f"Unrecognized accepted-variant record: {record_id}")
            for field, values in fields.items():
                if field not in FIELDS or not isinstance(values, list) or not all(
                    isinstance(value, str) for value in values
                ):
                    raise ValueError(f"Invalid accepted variants: {record_id} / {field}")
    role_pairs = {}
    use_record_rules = sheet_names[0].casefold() != "consensus"
    if use_record_rules:
        role_pairs = component("role_pair_variants", "role_pair_variants.json", {})
    if role_pairs:
        valid_slots = {slot for group in ROLE_GROUPS for slot in group}
        valid_ids = {row["record_id"] for row in records}
        if not isinstance(role_pairs, dict):
            raise ValueError("Role-pair variants must be a record-ID mapping.")
        for record_id, slots in role_pairs.items():
            if record_id not in valid_ids or not isinstance(slots, dict):
                raise ValueError(f"Invalid paired reference record: {record_id}")
            for slot, pairs in slots.items():
                if slot not in valid_slots or not isinstance(pairs, list):
                    raise ValueError(f"Invalid paired role: {record_id} / {slot}")
                if any(not isinstance(pair, dict) or "name" not in pair or
                       "amount_text" not in pair for pair in pairs):
                    raise ValueError(f"Each role variant needs a name and amount: {record_id} / {slot}")
    record_rules = {}
    if use_record_rules:
        record_rules = component("record_rules", "record_rules.json", {})
        validate_record_rules(record_rules, {row["record_id"] for row in records})
    return {
        "path": path, "records": records, "rejected": rejected, "papers": papers,
        "variants": variants, "companion_files": companion_files,
        "reference_sheet": sheet_names[0], "sheet_names": sheet_names,
        "role_pairs": role_pairs, "record_rules": record_rules,
        "embedded_rules": embedded,
    }


def save_reference_snapshot(reference, destination):
    """Save a complete reference that can be reused without companion files."""
    destination = Path(destination)
    if reference["embedded_rules"]:
        shutil.copyfile(reference["path"], destination)
    else:
        workbook = load_workbook(reference["path"])
        try:
            embed_reference_rules(workbook, {
                "papers": reference["papers"],
                "rejected_candidates": reference["rejected"],
                "accepted_variants": reference["variants"],
                "role_pair_variants": reference["role_pairs"],
                "record_rules": reference["record_rules"],
            })
            workbook.save(destination)
        finally:
            workbook.close()
    return destination


def prediction_content(row):
    """Canonical paired entities make IDs independent of reagent-slot order."""
    role_fields = {slot + suffix for group in ROLE_GROUPS for slot in group
                   for suffix in ("", "_amount_text")}
    content = {key: row[key] for key in ("doi", "mof_name") + FIELDS if key not in role_fields}
    content["roles"] = {
        group[0].split("_")[0]: sorted(
            (text(row.get(slot)), text(row.get(slot + "_amount_text"))) for slot in group)
        for group in ROLE_GROUPS
    }
    return content


def load_predictions(path, reference):
    headers, source = read_csv(path)
    if not ({"doi", "DOI"} & set(headers)):
        raise ValueError("The prediction CSV must contain doi or DOI.")
    if "mof_name" not in headers and not set(FIELDS).intersection(headers):
        raise ValueError("No synthesis fields found in the prediction CSV.")
    papers = {row["doi"]: row for row in reference["papers"]}
    predictions, ignored = [], []
    occurrences = Counter()
    for line, raw in enumerate(source, 2):
        row = projected_row(raw)
        doi = row["doi"]
        reason = ""
        if doi not in papers:
            reason = "outside_benchmark" if doi else "missing_doi"
        elif not papers[doi]["included"]:
            reason = "excluded_paper"
        elif text(raw.get("status")) and clean(raw.get("status")) != "ok":
            reason = "failed_or_invalid_status"
        elif not any(not is_absent(row[field]) for field in ("mof_name",) + FIELDS):
            reason = "empty_synthesis_placeholder"
        if reason:
            ignored.append({"csv_line": line, "doi": doi, "reason": reason,
                            "status": text(raw.get("status")), "error": text(raw.get("error"))})
            continue
        payload = json.dumps(prediction_content(row), sort_keys=True, ensure_ascii=False)
        identity = hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]
        occurrences[identity] += 1
        row["prediction_id"] = f"PRED-{identity}-{occurrences[identity]:03d}"
        row["csv_line"] = line
        row["parsed_json"] = text(raw.get("parsed_json"))
        row["paper_id"] = papers[doi]["paper_id"]
        predictions.append(row)
    return predictions, ignored, {
        "input_rows": len(source), "input_dois": len({doi_key(r.get("doi") or r.get("DOI")) for r in source}),
        "usable_prediction_records": len(predictions),
        "papers_with_usable_predictions": len({r["doi"] for r in predictions}),
        "ignored_rows_by_reason": dict(Counter(row["reason"] for row in ignored)),
        "missing_input_columns": [field for field in FIELDS if field not in headers],
    }


def name_key(value):
    return re.sub(r"[\W_]+", "", clean(value))


def amount_reference(slot, predicted_name, reference, variants=None, role_pairs=None):
    """Restrict amount alternatives to recorded variants of the matched name."""
    variants, role_pairs = variants or {}, role_pairs or {}
    field = slot + "_amount_text"
    pairs = role_pairs.get(slot, [])
    if not pairs:
        return reference.get(field, ""), variants.get(field, ()), "no_pair_constraints"
    matching_pairs = [
        pair for pair in pairs
        if value_matches(predicted_name, pair["name"], slot, pair.get("name_aliases", ()))
    ]
    if not matching_pairs:
        return reference.get(field, ""), variants.get(field, ()), "identity_unresolved"
    amounts = []
    for pair in matching_pairs:
        for value in [pair["amount_text"]] + list(pair.get("amount_aliases", ())):
            for alternative in split_alternatives(value):
                alternative = "not reported" if is_absent(alternative) else text(alternative)
                if alternative not in amounts:
                    amounts.append(alternative)
    return ALTERNATIVE_SEPARATOR.join(amounts), (), "paired_variant"


def validate_record_rules(rules, valid_ids):
    """Validate explicit record-level exceptions without evaluating text as code."""
    slots = {slot for group in ROLE_GROUPS for slot in group}
    if not isinstance(rules, dict) or set(rules) - valid_ids:
        raise ValueError("Record rules require recognized reference record IDs.")
    for record_id, settings in rules.items():
        if not isinstance(settings, dict) or set(settings) - {
            "role_reassignments", "distinct_components", "amount_pair_options",
            "composite_projections"
        }:
            raise ValueError(f"Invalid record rule: {record_id}")
        for rule in settings.get("role_reassignments", []):
            if (not isinstance(rule, dict) or not rule.get("names")
                    or not all(isinstance(name, str) for name in rule["names"])
                    or not rule.get("from_slots") or not rule.get("to_slots")
                    or set(rule["from_slots"] + rule["to_slots"]) - slots
                    or set(rule["from_slots"]) & set(rule["to_slots"])):
                raise ValueError(f"Invalid role reassignment: {record_id}")
        for rule in settings.get("composite_projections", []):
            if (not isinstance(rule, dict) or not rule.get("from_slots")
                    or set(rule["from_slots"]) - slots
                    or any(not rule.get(key) or any(
                        not isinstance(value, str) or is_absent(value)
                        for value in rule[key]) for key in ("source_names", "source_amounts"))
                    or len(rule.get("products", [])) < 2):
                raise ValueError(f"Invalid composite projection: {record_id}")
            for product in rule["products"]:
                if (not isinstance(product, dict) or not product.get("to_slots")
                        or not any(set(product["to_slots"]) <= set(group)
                                   for group in ROLE_GROUPS)
                        or any(not isinstance(product.get(key), str)
                               or is_absent(product[key])
                               for key in ("name", "amount_text"))):
                    raise ValueError(f"Invalid composite product: {record_id}")
        for group in settings.get("distinct_components", []):
            if (not isinstance(group, list) or len(group) < 2
                    or len(group) != len(set(group)) or set(group) - slots):
                raise ValueError(f"Invalid distinct-component rule: {record_id}")
        for rule in settings.get("amount_pair_options", []):
            if (not isinstance(rule, dict) or len(rule.get("slots", [])) != 2
                    or len(set(rule["slots"])) != 2 or set(rule["slots"]) - slots
                    or not rule.get("values_ml")
                    or any(len(pair) != 2 or any(
                        not isinstance(value, (int, float)) or not math.isfinite(value)
                        or value <= 0 for value in pair) for pair in rule["values_ml"])
                    or not 0 <= rule.get("absolute_tolerance_ml", 0) <= 0.01):
                raise ValueError(f"Invalid amount-pair rule: {record_id}")


def apply_role_reassignments(candidate, record_rules=None):
    """Apply exact recorded projections, then move complete permitted pairs."""
    candidate = dict(candidate)
    origins = {slot: slot for group in ROLE_GROUPS for slot in group}
    projected_slots = set()
    for rule in (record_rules or {}).get("composite_projections", []):
        source_names = {clean(value) for value in rule["source_names"]}
        source_amounts = {clean(value) for value in rule["source_amounts"]}
        sources = sorted(rule["from_slots"], key=lambda slot: (
            clean(candidate.get(slot)), clean(candidate.get(slot + "_amount_text"))))
        for source in sources:
            # Require both complete literal expressions. Quantity parsing or
            # chemical aliases must not authorize an unlisted decomposition.
            if (clean(candidate.get(source)) not in source_names
                    or clean(candidate.get(source + "_amount_text")) not in source_amounts):
                continue
            trial = dict(candidate)
            trial[source] = trial[source + "_amount_text"] = ""
            destinations = []
            for product in rule["products"]:
                available = product["to_slots"]
                if any(value_matches(trial.get(slot), product["name"], slot)
                       for slot in available):
                    break
                destination = next((slot for slot in available
                                    if is_absent(trial.get(slot))
                                    and is_absent(trial.get(slot + "_amount_text"))), None)
                if destination is None:
                    break
                trial[destination] = product["name"]
                trial[destination + "_amount_text"] = product["amount_text"]
                destinations.append(destination)
            else:
                candidate = trial
                for destination in destinations:
                    origins[destination] = origins[source]
                projected_slots.update(destinations)
                break
    for rule in (record_rules or {}).get("role_reassignments", []):
        def allowed(slot):
            return any(value_matches(candidate.get(slot), name, slot)
                       for name in rule["names"])
        # A second copy is never silently removed or merged with the first.
        if any(allowed(slot) for slot in rule["to_slots"]):
            continue
        sources = sorted(rule["from_slots"], key=lambda slot: (
            canonical_name(candidate.get(slot)), clean(candidate.get(slot + "_amount_text"))))
        for source in sources:
            if not allowed(source):
                continue
            destination = next((slot for slot in rule["to_slots"]
                                if is_absent(candidate.get(slot))
                                and is_absent(candidate.get(slot + "_amount_text"))), None)
            if destination is None:
                break
            for suffix in ("", "_amount_text"):
                candidate[destination + suffix], candidate[source + suffix] = (
                    candidate.get(source + suffix, ""), candidate.get(destination + suffix, ""))
            origins[destination], origins[source] = origins[source], origins[destination]
            if source in projected_slots:
                projected_slots.remove(source)
                projected_slots.add(destination)
            break
    candidate["_composite_projection_slots"] = tuple(sorted(projected_slots))
    return candidate, origins


def rule_field_matches(candidate, reference, variants=None, record_rules=None):
    """Return explicit true/false field decisions from coupled reference rules."""
    variants, record_rules = variants or {}, record_rules or {}
    decisions = {}
    for group in record_rules.get("distinct_components", []):
        seen = []
        for slot in group:
            name = candidate.get(slot)
            if is_absent(name):
                continue
            if any(value_matches(name, previous, slot) for previous in seen):
                decisions[slot] = (False, "duplicate_component")
            seen.append(name)
    for rule in record_rules.get("amount_pair_options", []):
        slots = rule["slots"]
        if not all(value_matches(candidate.get(slot), reference.get(slot), slot,
                                 variants.get(slot, ())) for slot in slots):
            continue
        quantities = [_quantities(candidate.get(slot + "_amount_text"), slot + "_amount_text")
                      for slot in slots]
        if not all(set(value) == {"ml"} and len(value["ml"]) == 1 for value in quantities):
            continue
        values = [value["ml"][0] for value in quantities]
        tolerance = rule.get("absolute_tolerance_ml", 0)
        if any(all(math.isclose(a, b, rel_tol=1e-6, abs_tol=tolerance)
                   for a, b in zip(values, pair))
               for option in rule["values_ml"] for pair in (option, list(reversed(option)))):
            decisions.update({slot + "_amount_text": (True, "accepted_volume_pair")
                              for slot in slots})
    return decisions


def align_role_slots(candidate, reference, variants=None, role_pairs=None, record_rules=None):
    """Pair reagent entities within a role, keeping each name and amount together.

    A lexicographic objective first maximizes matching nonempty identities,
    then accepted absence, nonempty amount matches, and presence agreement.
    Content provides the final tie-break, independent of input slot order.
    At most six permutations are considered for a three-slot role.
    """
    variants = variants or {}
    candidate, origins = apply_role_reassignments(candidate, record_rules)
    aligned = dict(candidate)
    aligned["_composite_projection_slots"] = []
    slot_map = {}
    for group in ROLE_GROUPS:
        best = None
        best_content = None
        best_permutation = group
        for permutation in itertools.permutations(group):
            name_matches = absence_matches = amount_matches = presence_matches = 0
            for reference_slot, candidate_slot in zip(group, permutation):
                predicted_name = candidate.get(candidate_slot)
                expected_name = reference.get(reference_slot)
                if not is_unknown(expected_name):
                    presence_matches += bool(not is_absent(predicted_name)) == bool(
                        alternatives(expected_name, variants.get(reference_slot, ())))
                    if is_absent(predicted_name):
                        absence_matches += accepts_absence(
                            expected_name, variants.get(reference_slot, ()))
                    else:
                        name_matches += value_matches(
                            predicted_name, expected_name, reference_slot,
                            variants.get(reference_slot, ()))
                amount_field = reference_slot + "_amount_text"
                predicted_amount = candidate.get(candidate_slot + "_amount_text")
                expected_amount, amount_aliases, _ = amount_reference(
                    reference_slot, predicted_name, reference, variants, role_pairs)
                if not is_absent(predicted_amount) and not is_unknown(expected_amount):
                    amount_matches += value_matches(
                        predicted_amount, expected_amount, amount_field, amount_aliases)
            key = (name_matches, absence_matches, amount_matches, presence_matches)
            content = tuple((is_absent(candidate.get(slot)), canonical_name(candidate.get(slot)),
                             clean(candidate.get(slot + "_amount_text"))) for slot in permutation)
            if best is None or key > best or (key == best and content < best_content):
                best, best_content, best_permutation = key, content, permutation
        for reference_slot, candidate_slot in zip(group, best_permutation):
            aligned[reference_slot] = candidate.get(candidate_slot, "")
            aligned[reference_slot + "_amount_text"] = candidate.get(candidate_slot + "_amount_text", "")
            slot_map[reference_slot] = origins[candidate_slot]
            if candidate_slot in candidate.get("_composite_projection_slots", ()):
                aligned["_composite_projection_slots"].append(reference_slot)
    return aligned, slot_map


def pair_score(candidate, reference, variants=None, role_pairs=None, record_rules=None):
    """Score record alignment independently of exact field-value scoring."""
    variants = variants or {}
    candidate, _ = align_role_slots(candidate, reference, variants, role_pairs, record_rules)
    rule_matches = rule_field_matches(candidate, reference, variants, record_rules)
    left, right = name_key(candidate.get("mof_name")), name_key(reference.get("mof_name"))
    name_similarity = (difflib.SequenceMatcher(None, left, right).ratio()
                       if left and right else None)
    compared = agreed = 0
    for field in MATCH_ID_FIELDS:
        expected = alternatives(reference.get(field), variants.get(field, ()))
        value = candidate.get(field)
        if is_absent(value) or not expected or is_unknown(reference.get(field)):
            continue
        compared += 1
        normalized = re.sub(r"\s+", "", canonical_name(value))
        same = value_matches(value, reference.get(field), field, variants.get(field, ()))
        # Substring similarity is confined to alignment. It never awards a
        # field-level TP, and salts/hydrates remain distinct in final scoring.
        same = same or any(
            min(len(normalized), len(other)) >= 4
            and (normalized in other or other in normalized)
            for other in (re.sub(r"\s+", "", canonical_name(v)) for v in expected)
        )
        agreed += rule_matches.get(field, (bool(same), ""))[0]
    for field in MATCH_AMOUNT_FIELDS:
        value = candidate.get(field)
        expected_value, extra = reference.get(field), variants.get(field, ())
        if field.endswith("_amount_text"):
            slot = field.removesuffix("_amount_text")
            expected_value, extra, _ = amount_reference(
                slot, candidate.get(slot), reference, variants, role_pairs)
        expected = alternatives(expected_value, extra)
        if field in rule_matches and rule_matches[field][0]:
            compared += 1
            agreed += 1
            continue
        if is_absent(value) or not expected or not _quantities(value, field):
            continue
        comparable = [v for v in expected if _quantities(value, field).keys() &
                      _quantities(v, field).keys()]
        if not comparable:
            continue
        compared += 1
        agreed += any(_quantity_equal(value, v, field, relative_tolerance=0.02)
                      for v in comparable)
    field_similarity = agreed / compared if compared else None
    if name_similarity is None and field_similarity is None:
        score = 0.0
    elif name_similarity is None:
        score = field_similarity * (0.9 if compared >= 3 else 0.6)
    elif field_similarity is None:
        score = 0.8 * name_similarity
    else:
        score = 0.35 * name_similarity + 0.65 * field_similarity
    # Only populated fields contribute to the tie-break. Shared blanks cannot
    # dominate the alignment of sparsely described recipes.
    tie_fields = [
        field for field in FIELDS
        if not is_unknown(reference.get(field))
        and (not is_absent(candidate.get(field)) or
             alternatives(reference.get(field), variants.get(field, ())))
    ]
    tie_matches = 0
    for field in tie_fields:
        expected_value, extra = reference.get(field), variants.get(field, ())
        if field.endswith("_amount_text"):
            slot = field.removesuffix("_amount_text")
            expected_value, extra, _ = amount_reference(
                slot, candidate.get(slot), reference, variants, role_pairs)
        tie_matches += rule_matches.get(field, (
            value_matches(candidate.get(field), expected_value, field, extra), ""))[0]
    tie = tie_matches / len(tie_fields) if tie_fields else 0
    return float(score), float(tie), compared


def match_records(predictions, reference, minimum=0.50):
    """Find one-to-one matches within each DOI with explicit unmatched choices."""
    if not 0 <= minimum <= 1:
        raise ValueError("The matching threshold must be between zero and one.")
    included = {row["doi"] for row in reference["papers"] if row["included"]}
    candidates_by_doi, targets_by_doi = defaultdict(list), defaultdict(list)
    for row in predictions:
        candidates_by_doi[row["doi"]].append(row)
    for kind in ("records", "rejected"):
        for row in reference[kind]:
            if row["doi"] in included:
                targets_by_doi[row["doi"]].append((kind, row))
    assignments, diagnostics = {}, []
    for doi in sorted(included):
        candidates = sorted(candidates_by_doi[doi], key=lambda row: row["prediction_id"])
        targets = sorted(targets_by_doi[doi], key=lambda pair: pair[1]["record_id"])
        if not candidates:
            continue
        scores = np.zeros((len(candidates), len(targets)))
        ties, comparisons = np.zeros_like(scores), np.zeros_like(scores)
        for i, candidate in enumerate(candidates):
            for j, (_, target) in enumerate(targets):
                scores[i, j], ties[i, j], comparisons[i, j] = pair_score(
                    candidate, target, reference["variants"].get(target["record_id"], {}),
                    reference.get("role_pairs", {}).get(target["record_id"], {}),
                    reference.get("record_rules", {}).get(target["record_id"], {})
                )
        objective = np.zeros((len(candidates), len(targets) + len(candidates)))
        if targets:
            objective[:, :len(targets)] = np.where(
                scores >= minimum, scores + 0.03 * ties, -1e6)
        row_indices, column_indices = linear_sum_assignment(-objective)
        for i, j in zip(row_indices, column_indices):
            candidate = candidates[i]
            selected = j < len(targets)
            if selected:
                kind, target = targets[j]
                assignments[candidate["prediction_id"]] = {
                    "kind": kind, "record_id": target["record_id"],
                    "score": float(scores[i, j]),
                }
            order = np.argsort(-scores[i], kind="stable") if targets else []
            best = int(order[0]) if len(order) else None
            second = int(order[1]) if len(order) > 1 else None
            chosen_score = float(scores[i, j]) if selected else None
            competing = np.delete(scores[i], j) if selected else scores[i]
            margin = chosen_score - float(competing.max()) if selected and len(competing) else None
            diagnostics.append({
                "doi": doi, "prediction_id": candidate["prediction_id"],
                "csv_line": candidate["csv_line"],
                "assigned_record_id": targets[j][1]["record_id"] if selected else "",
                "assigned_kind": targets[j][0] if selected else "unmatched",
                "match_score": chosen_score,
                "recipe_fields_compared": int(comparisons[i, j]) if selected else 0,
                "best_record_id": targets[best][1]["record_id"] if best is not None else "",
                "best_score": float(scores[i, best]) if best is not None else None,
                "second_best_score": float(scores[i, second]) if second is not None else None,
                "assigned_margin": margin,
                "ambiguous_alignment": bool(selected and margin is not None and margin <= 0.03),
                "name_only_alignment": bool(selected and comparisons[i, j] == 0),
            })
    return assignments, diagnostics


def divide(numerator, denominator):
    return numerator / denominator if denominator else None


def confusion_metrics(tp, fp, tn, fn):
    """Binary metrics; undefined values stay null rather than silently zero."""
    precision, recall = divide(tp, tp + fp), divide(tp, tp + fn)
    specificity = divide(tn, tn + fp) if tn is not None else None
    return {
        "TP": tp, "FP": fp, "TN": tn, "FN": fn,
        "accuracy": divide(tp + tn, tp + tn + fp + fn) if tn is not None else None,
        "balanced_accuracy": ((recall + specificity) / 2
                              if recall is not None and specificity is not None else None),
        "precision": precision, "recall": recall,
        "F1": divide(2 * tp, 2 * tp + fp + fn),
    }


def field_label(candidate, reference, field, additional=(), alignment="matched", override_match=None):
    """Label a single aligned field and expose both value and presence counts."""
    counts = {key: 0 for key in ("TP", "FP", "TN", "FN", "presence_TP",
                                "presence_FP", "presence_TN", "presence_FN")}
    if is_unknown(reference):
        return {"label": "unresolved_reference", "assessable": False,
                "exact_correct": None, "optional_reference": False, **counts}
    expected = alternatives(reference, additional)
    allow_absence = accepts_absence(reference, additional)
    optional = bool(expected and allow_absence)
    present = not is_absent(candidate)
    if alignment == "missed" and (not expected or optional):
        return {"label": ("optional_field_on_missed_record" if optional else
                          "absent_field_on_missed_record"), "assessable": False,
                "exact_correct": None, "optional_reference": optional, **counts}
    if alignment == "extra" and not present:
        return {"label": "absent_field_on_extra_record", "assessable": False,
                "exact_correct": None, "optional_reference": False, **counts}
    if present and expected:
        counts["presence_TP"] = int(not optional)
        correct = (value_matches(candidate, reference, field, additional)
                   if override_match is None else override_match)
        if correct:
            counts["TP"] = 1
            label = "correct_value"
        else:
            counts["FP"] = 1
            counts["FN"] = int(not optional)
            label = "wrong_value"
    elif expected and not allow_absence:
        counts["FN"] = counts["presence_FN"] = 1
        correct, label = False, "missing_value"
    elif present:
        counts["FP"] = counts["presence_FP"] = 1
        correct, label = False, "extra_value"
    else:
        counts["TN"] = 1
        counts["presence_TN"] = int(not optional)
        correct, label = True, "correct_absence"
    return {"label": label, "assessable": True, "exact_correct": correct,
            "optional_reference": optional, **counts}


def label_records(predictions, reference, assignments):
    included = {row["doi"] for row in reference["papers"] if row["included"]}
    prediction_map = {row["prediction_id"]: row for row in predictions}
    found = {match["record_id"]: (prediction_map[pid], match)
             for pid, match in assignments.items() if match["kind"] == "records"}
    record_labels, field_labels = [], []
    for expected in reference["records"]:
        if expected["doi"] not in included:
            continue
        record_id = expected["record_id"]
        candidate, match = found.get(record_id, ({}, {}))
        alignment = "matched" if candidate else "missed"
        aligned_candidate, slot_map = align_role_slots(
            candidate, expected, reference["variants"].get(record_id, {}),
            reference.get("role_pairs", {}).get(record_id, {}),
            reference.get("record_rules", {}).get(record_id, {}))
        rule_matches = rule_field_matches(
            aligned_candidate, expected, reference["variants"].get(record_id, {}),
            reference.get("record_rules", {}).get(record_id, {}))
        record = {key: expected[key] for key in IDENTIFIERS}
        record.update({
            "prediction_id": candidate.get("prediction_id", ""),
            "csv_line": candidate.get("csv_line"),
            "record_label": "TP" if candidate else "FN",
            "match_score": match.get("score"),
        })
        these_fields = []
        for field in FIELDS:
            additional = reference["variants"].get(record_id, {}).get(field, ())
            effective_reference = expected.get(field)
            pair_status = ""
            if field.endswith("_amount_text") and candidate:
                slot = field.removesuffix("_amount_text")
                effective_reference, additional, pair_status = amount_reference(
                    slot, aligned_candidate.get(slot), expected,
                    reference["variants"].get(record_id, {}),
                    reference.get("role_pairs", {}).get(record_id, {}))
            label = field_label(aligned_candidate.get(field), effective_reference, field,
                                additional, alignment, rule_matches.get(field, (None, ""))[0])
            reference_slot = field.removesuffix("_amount_text") if field.startswith(
                ("metal_", "linker_", "modulator_", "solvent_")) else ""
            prediction_slot = slot_map.get(reference_slot, "") if candidate else ""
            source_field = (prediction_slot + ("_amount_text" if field.endswith("_amount_text") else "")
                            if prediction_slot else field if candidate else "")
            applied_rule = rule_matches.get(field, (None, ""))[1]
            if reference_slot in aligned_candidate.get("_composite_projection_slots", ()):
                applied_rule = applied_rule or "explicit_composite_projection"
            if (prediction_slot and reference_slot
                    and prediction_slot.split("_")[0] != reference_slot.split("_")[0]):
                applied_rule = applied_rule or "permitted_role_reassignment"
            entry = {
                "record_id": record_id, "prediction_id": candidate.get("prediction_id", ""),
                "paper_id": expected["paper_id"], "doi": expected["doi"],
                "alignment": alignment, "field": field,
                "reference_value": expected.get(field, ""),
                "effective_reference_value": effective_reference,
                "accepted_variants": ALTERNATIVE_SEPARATOR.join(additional),
                "prediction_value": aligned_candidate.get(field, ""),
                "reference_slot": reference_slot,
                "original_prediction_slot": prediction_slot,
                "original_prediction_field": source_field,
                "original_prediction_value": candidate.get(source_field, ""),
                "pair_identity_status": pair_status,
                "record_rule": applied_rule, **label,
            }
            field_labels.append(entry)
            these_fields.append(entry)
        record["assessable_fields"] = sum(row["assessable"] for row in these_fields)
        record["correct_fields"] = sum(row["exact_correct"] is True for row in these_fields)
        record["unresolved_fields"] = sum(row["label"] == "unresolved_reference" for row in these_fields)
        record["all_evaluated_fields_correct"] = bool(
            candidate and record["assessable_fields"] > 0
            and record["correct_fields"] == record["assessable_fields"]
        )
        record["all_26_fields_correct"] = bool(
            record["all_evaluated_fields_correct"] and record["assessable_fields"] == len(FIELDS)
        )
        record_labels.append(record)
    for candidate in predictions:
        pid = candidate["prediction_id"]
        match = assignments.get(pid, {})
        if match.get("kind") == "records":
            continue
        record_labels.append({
            "record_id": "", "paper_id": candidate["paper_id"], "doi": candidate["doi"],
            "mof_name": candidate["mof_name"], "prediction_id": pid,
            "csv_line": candidate["csv_line"], "record_label": "FP",
            "match_score": match.get("score"), "assessable_fields": None,
            "correct_fields": None, "unresolved_fields": 0,
            "all_evaluated_fields_correct": False, "all_26_fields_correct": False,
            "extra_reason": "known_rejected_candidate" if match else "unmatched_prediction",
            "rejected_candidate_id": match.get("record_id", ""),
        })
        # Extra records have no reference slot order. Canonical content order
        # keeps their per-field FP allocation independent of raw input slots.
        aligned_extra, extra_slots = align_role_slots(candidate, {})
        for field in FIELDS:
            label = field_label(aligned_extra[field], "", field, alignment="extra")
            slot = field.removesuffix("_amount_text") if field.startswith(
                ("metal_", "linker_", "modulator_", "solvent_")) else ""
            source_slot = extra_slots.get(slot, "")
            source_field = source_slot + ("_amount_text" if field.endswith("_amount_text") else "") if slot else field
            field_labels.append({
                "record_id": "", "prediction_id": pid, "paper_id": candidate["paper_id"],
                "doi": candidate["doi"], "alignment": "extra", "field": field,
                "reference_value": "", "accepted_variants": "",
                "effective_reference_value": "",
                "prediction_value": aligned_extra[field],
                "reference_slot": "",
                "original_prediction_slot": source_slot,
                "original_prediction_field": source_field, "pair_identity_status": "",
                "original_prediction_value": candidate.get(source_field, ""),
                "record_rule": "", **label,
            })
    return record_labels, field_labels


def summarize_fields(rows):
    assessable = [row for row in rows if row["assessable"]]
    counts = {key: sum(row[key] for row in assessable) for key in ("TP", "FP", "TN", "FN")}
    precision = divide(counts["TP"], counts["TP"] + counts["FP"])
    recall = divide(counts["TP"], counts["TP"] + counts["FN"])
    presence = confusion_metrics(
        *(sum(row["presence_" + key] for row in assessable) for key in ("TP", "FP", "TN", "FN"))
    )
    return {
        **counts, "precision": precision, "recall": recall,
        "F1": divide(2 * counts["TP"], 2 * counts["TP"] + counts["FP"] + counts["FN"]),
        "exact_cell_accuracy": divide(sum(row["exact_correct"] is True for row in assessable),
                                      len(assessable)),
        "assessable_cells": len(assessable),
        "unresolved_cells": sum(row["label"] == "unresolved_reference" for row in rows),
        "wrong_values": sum(row["label"] == "wrong_value" for row in rows),
        "optional_reference_cells": sum(row["optional_reference"] for row in rows),
        "optional_assessable_cells": sum(row["optional_reference"] for row in assessable),
        "presence_assessable_cells": sum(not row["optional_reference"] for row in assessable),
        "required_value_TP": sum(row["TP"] for row in assessable if not row["optional_reference"]),
        "required_value_FN": sum(row["FN"] for row in assessable if not row["optional_reference"]),
        "required_value_recall": divide(
            sum(row["TP"] for row in assessable if not row["optional_reference"]),
            sum(row["TP"] + row["FN"] for row in assessable if not row["optional_reference"])),
        **{"presence_" + key: value for key, value in presence.items()},
    }


def mean_defined(values):
    values = [value for value in values if value is not None]
    return sum(values) / len(values) if values else None


def summarize_field_groups(field_labels):
    """Pool slot-level counts and cell denominators before computing metrics."""
    by_field = defaultdict(list)
    for row in field_labels:
        by_field[row["field"]].append(row)
    metrics = []
    for scope in SCOPE_LABELS:
        for field, members in FIELD_GROUPS.items():
            rows = [
                row for member in members for row in by_field[member]
                if scope == "all_records" or row["alignment"] == "matched"
            ]
            metrics.append({"scope": scope, "field": field, **summarize_fields(rows)})
    return metrics


def summarize(record_labels, field_labels, reference, coverage):
    counts = Counter(row["record_label"] for row in record_labels)
    record_metrics = confusion_metrics(counts["TP"], counts["FP"], None, counts["FN"])
    record_metrics.update({
        "averaging": "micro_records",
        "reference_records": counts["TP"] + counts["FN"],
        "prediction_records": counts["TP"] + counts["FP"],
        "exact_records_26_fields": sum(row["all_26_fields_correct"] for row in record_labels),
        "exact_record_recall_26_fields": divide(
            sum(row["all_26_fields_correct"] for row in record_labels), counts["TP"] + counts["FN"]),
        "exact_record_accuracy_matched_26_fields": divide(
            sum(row["all_26_fields_correct"] for row in record_labels), counts["TP"]),
    })
    field_metrics = summarize_field_groups(field_labels)
    paper_rows = []
    records_by_doi, fields_by_doi = defaultdict(list), defaultdict(list)
    for row in record_labels:
        records_by_doi[row["doi"]].append(row)
    for row in field_labels:
        fields_by_doi[row["doi"]].append(row)
    for paper in reference["papers"]:
        if not paper["included"]:
            continue
        labels = records_by_doi[paper["doi"]]
        count = Counter(row["record_label"] for row in labels)
        tp, fp, fn = count["TP"], count["FP"], count["FN"]
        paper_rows.append({
            "paper_id": paper["paper_id"], "doi": paper["doi"],
            "reference_records": tp + fn, "recovered_records": tp,
            "predicted_records": tp + fp, "extra_records": fp,
            **confusion_metrics(tp, fp, None, fn),
            "fully_recalled": fn == 0, "no_extra_records": fp == 0,
            "record_set_correct": fn == 0 and fp == 0,
            "field_precision": summarize_fields(fields_by_doi[paper["doi"]])["precision"],
            "field_recall": summarize_fields(fields_by_doi[paper["doi"]])["recall"],
        })
    scopes = {}
    for scope in SCOPE_LABELS:
        rows = field_labels if scope == "all_records" else [
            row for row in field_labels if row["alignment"] == "matched"]
        fields = [row for row in field_metrics if row["scope"] == scope]
        scopes[scope] = {
            "micro": summarize_fields(rows),
            "macro_fields": {
                metric: mean_defined(row[metric] for row in fields)
                for metric in ("precision", "recall", "F1", "exact_cell_accuracy",
                               "presence_accuracy", "presence_balanced_accuracy")
            },
            "macro_fields_defined_counts": {
                metric: sum(row[metric] is not None for row in fields)
                for metric in ("precision", "recall", "F1", "exact_cell_accuracy",
                               "presence_accuracy", "presence_balanced_accuracy")
            },
        }
    summary = {
        "schema_version": SCHEMA_VERSION, "reference_sheet": reference.get("reference_sheet", SHEETS[0]),
        "record_detection": record_metrics,
        "field_evaluation": scopes,
        "field_groups": {field: list(members) for field, members in FIELD_GROUPS.items()},
        "scope_labels": dict(SCOPE_LABELS),
        "paper_evaluation": {
            "sampled_papers": len(reference["papers"]), "evaluated_papers": len(paper_rows),
            "excluded_papers": sum(not row["included"] for row in reference["papers"]),
            "mean_reference_syntheses": mean_defined(row["reference_records"] for row in paper_rows),
            "mean_recovered_syntheses": mean_defined(row["recovered_records"] for row in paper_rows),
            "median_reference_syntheses": float(np.median([r["reference_records"] for r in paper_rows])) if paper_rows else None,
            "median_recovered_syntheses": float(np.median([r["recovered_records"] for r in paper_rows])) if paper_rows else None,
            "macro_paper_precision": mean_defined(row["precision"] for row in paper_rows),
            "macro_paper_recall": mean_defined(row["recall"] for row in paper_rows),
            "macro_paper_F1": mean_defined(row["F1"] for row in paper_rows),
            "macro_paper_precision_defined_n": sum(row["precision"] is not None for row in paper_rows),
            "macro_paper_recall_defined_n": sum(row["recall"] is not None for row in paper_rows),
            "fully_recalled_papers": sum(row["fully_recalled"] for row in paper_rows),
            "papers_without_extra_records": sum(row["no_extra_records"] for row in paper_rows),
            "record_set_correct_papers": sum(row["record_set_correct"] for row in paper_rows),
        },
        "coverage": coverage,
        "definitions": {
            "record_TN": "Undefined: no finite universe of nonsynthesis records is sampled.",
            "record_accuracy_balanced_accuracy": "Undefined because record TN is not defined.",
            "wrong_nonempty_value": "One FP for the predicted value and one FN for the reference value.",
            "optional_reference": "Where both absence and a nonempty value are accepted, absence is correct; "
                                  "a wrong nonempty value is FP only; a missed record contributes no field FN "
                                  "for that optional field. Record FN remains unchanged.",
            "optional_presence": "Optional-reference cells are excluded from binary presence metrics.",
            "required_value_recall": "Recall restricted to nonempty fields whose accepted values do not include absence.",
            "field_TN": "Matched record with both reference and prediction absent.",
            "field_accuracy": "Use exact_cell_accuracy; FP+FN for a wrong value must not double its denominator.",
            "presence_metrics": "Binary present/absent classification; a wrong nonempty value is presence TP.",
            "all_records": "All-protocol evaluation includes matched records and populated fields on "
                           "missed reference records and extra prediction records.",
            "matched_records": "Recovered-protocol evaluation includes only fields on matched records.",
            "empty_unmatched_fields": "Do not contribute TN because no corresponding record exists.",
            "pooled_fields": "Sum counts and cell denominators across the listed reagent slots before "
                             "computing each of the 14 field-type metrics; names and amounts are separate.",
            "macro_fields": "Unweighted mean of defined metrics across the 14 pooled field types; "
                            "defined counts are reported.",
            "macro_papers": "Unweighted mean of defined per-paper metrics; no-output paper recall is zero when reference positive.",
            "micro": "Pool numerator and denominator counts before computing the metric.",
            "reference_scope": "Recovery relative to the fixed reviewed reference, not an exhaustive census of all literature syntheses.",
        },
    }
    return summary, field_metrics, paper_rows


def bootstrap_papers(record_labels, field_labels, paper_rows, repetitions=2000, seed=42):
    """Percentile intervals resample papers, keeping their records clustered."""
    if repetitions < 0:
        raise ValueError("Bootstrap repetitions must be nonnegative.")
    if repetitions == 0 or not paper_rows:
        return {"repetitions": repetitions, "seed": seed, "unit": "paper", "intervals": {}}
    dois = [row["doi"] for row in paper_rows]
    index = {doi: i for i, doi in enumerate(dois)}
    counts = np.zeros((len(dois), 6), dtype=float)
    for row in record_labels:
        counts[index[row["doi"]], {"TP": 0, "FP": 1, "FN": 2}[row["record_label"]]] += 1
    for row in field_labels:
        for column, label in enumerate(("TP", "FP", "FN"), 3):
            counts[index[row["doi"]], column] += row[label]
    generator = np.random.default_rng(seed)
    samples = defaultdict(list)
    for _ in range(repetitions):
        total = counts[generator.integers(0, len(dois), size=len(dois))].sum(axis=0)
        for offset, level in ((0, "record"), (3, "field")):
            tp, fp, fn = total[offset:offset + 3]
            for name, value in (("precision", divide(tp, tp + fp)),
                                ("recall", divide(tp, tp + fn)),
                                ("F1", divide(2 * tp, 2 * tp + fp + fn))):
                if value is not None:
                    samples[f"{level}_{name}"].append(value)
    return {
        "repetitions": repetitions, "seed": seed, "unit": "paper",
        "method": "95% percentile interval; papers resampled with replacement",
        "intervals": {key: {"low": float(np.percentile(values, 2.5)),
                            "high": float(np.percentile(values, 97.5)),
                            "defined_replicates": len(values)}
                      for key, values in samples.items()},
    }


def evaluate(csv_path=DEFAULT_CSV, ground_truth=DEFAULT_REFERENCE, *,
             output_dir=None, output_root=None,
             match_threshold=0.50, bootstraps=2000, seed=42):
    """Validate, align, label, and report a saved extraction without API access."""
    csv_path = Path(csv_path).expanduser().resolve()
    reference = load_reference(ground_truth)
    predictions, ignored, coverage = load_predictions(csv_path, reference)
    assignments, matching = match_records(predictions, reference, match_threshold)
    record_labels, field_labels = label_records(predictions, reference, assignments)
    summary, field_metrics, paper_rows = summarize(
        record_labels, field_labels, reference, coverage)
    summary["confidence_intervals"] = bootstrap_papers(
        record_labels, field_labels, paper_rows, bootstraps, seed)
    summary["matching"] = {
        "threshold": match_threshold,
        "ambiguous_alignments": sum(row["ambiguous_alignment"] for row in matching),
        "name_only_alignments": sum(row["name_only_alignment"] for row in matching),
    }
    if output_dir:
        destination = Path(output_dir).expanduser().resolve()
        destination.mkdir(parents=True, exist_ok=False)
    else:
        destination = create_run_directory(
            output_root or DEFAULT_OUTPUT_ROOT)
    save_csv(destination / "record_labels.csv", record_labels)
    save_csv(destination / "field_labels.csv", field_labels)
    save_csv(destination / "field_metrics.csv", field_metrics)
    save_csv(destination / "paper_metrics.csv", paper_rows)
    save_csv(destination / "matching_diagnostics.csv", matching)
    save_csv(destination / "matching_uncertain.csv",
             [row for row in matching if row["ambiguous_alignment"] or row["name_only_alignment"]],
             list(matching[0]) if matching else ["prediction_id", "assigned_record_id", "match_score"])
    save_csv(destination / "ignored_input_rows.csv", ignored,
             ["csv_line", "doi", "reason", "status", "error"])
    save_csv(destination / "predictions_used.csv", predictions)
    save_csv(destination / "reference_used.csv", reference["records"], IDENTIFIERS + FIELDS)
    save_csv(destination / "paper_manifest_used.csv", reference["papers"])
    snapshot = save_reference_snapshot(reference, destination / "reference_workbook_used.xlsx")
    (destination / "summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False, allow_nan=False), encoding="utf-8")
    manifest = {
        "schema_version": SCHEMA_VERSION, "created_utc": datetime.now(timezone.utc).isoformat(),
        "completed": True,
        "input_csv": {"path": str(csv_path), "sha256": digest(csv_path)},
        "reference_workbook": {"path": str(reference["path"]), "sha256": digest(reference["path"])},
        "reference_snapshot": {"path": snapshot.name, "sha256": digest(snapshot)},
        "reference_companions": {
            key: {"path": str(path), "sha256": digest(path)}
            for key, path in reference["companion_files"].items()
        },
        "implementation_sha256": digest(__file__),
        "settings": {"match_threshold": match_threshold, "bootstraps": bootstraps, "seed": seed},
        "evaluated_fields": list(FIELDS), "sheet_names": list(reference["sheet_names"]),
        "field_groups": summary["field_groups"], "scope_labels": summary["scope_labels"],
        "matching_protocol": {
            "within_doi": True, "one_to_one": True, "algorithm": "Hungarian assignment with unmatched choices",
            "role_slots": "Within each metal/linker/modulator/solvent role, names and amounts move together "
                          "under the identity-first optimal permutation; content resolves ties independently "
                          "of input slots. Extra records use canonical content order. Entities cannot be reused.",
            "prediction_identity": "Hash of DOI, MOF name, scalar fields and sorted name/amount entities per role.",
            "paired_amount_variants": "Where recorded name/amount pairs exist and a predicted name matches, "
                                     "only amounts attached to that name are accepted; unresolved identity is flagged.",
            "record_rules": "Only explicitly listed records allow cross-role placement or named solvent-volume "
                            "pair alternatives. Distinct-component constraints apply only where listed. "
                            "Names and amounts remain attached during cross-role movement. Composite projections "
                            "require both listed complete source expressions and empty destinations, preserve "
                            "original values in labels, and never infer quantities or duplicate existing components.",
            "name_weight": 0.35, "recipe_weight": 0.65,
            "alignment_numeric_relative_tolerance": 0.02, "field_value_relative_tolerance": 1e-6,
            "nonempty_field_tiebreak_weight": 0.03,
            "rejected_candidates": "Competing alignment targets whose predictions count as FP, never TN.",
            "field_schema": "26 semantic fields with explicit unit normalization and stable record identifiers.",
        },
    }
    (destination / "run_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    return destination, summary


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("csv", nargs="?", type=Path, default=DEFAULT_CSV,
                        help="Raw extraction CSV; defaults to the bundled 200-paper extraction.")
    parser.add_argument("--ground-truth", "--reference", type=Path, default=DEFAULT_REFERENCE,
                        help="Reference workbook with embedded matching metadata.")
    parser.add_argument("--output-dir", type=Path, help="New destination; existing folders are refused.")
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT,
                        help="Parent for numbered runs; defaults to results/positive_extraction.")
    parser.add_argument("--match-threshold", type=float, default=0.50)
    parser.add_argument("--bootstraps", type=int, default=2000)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--validate-only", action="store_true",
                        help="Validate the reference and input without creating output.")
    args = parser.parse_args(argv)
    if args.validate_only:
        reference = load_reference(args.ground_truth)
        _, _, coverage = load_predictions(args.csv, reference)
        print(json.dumps({"reference_records": len(reference["records"]),
                          "sampled_papers": len(reference["papers"]), **coverage}, indent=2))
        return 0
    destination, summary = evaluate(
        args.csv, args.ground_truth, output_dir=args.output_dir, output_root=args.output_root,
        match_threshold=args.match_threshold, bootstraps=args.bootstraps, seed=args.seed)
    record = summary["record_detection"]
    print(f"Saved evaluation: {display_path(destination)}")
    print(f"Record detection: TP={record['TP']}, FP={record['FP']}, FN={record['FN']}; TN undefined.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
