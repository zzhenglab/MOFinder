"""Analyse MOF Quest responses by their recorded reaction identifiers."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import json
from math import sqrt
from pathlib import Path
from statistics import mean, stdev
from typing import Any


RESPONSE_LABELS = {
    "Very Confident Success": "P",
    "Likely Success": "P",
    "Likely Fail": "N",
    "Very Confident Fail": "N",
}
EXPERIENCE_ORDER = ("<1 year", "1-3 years", ">3 years")
QUESTION_FIELDS = ("question", "reaction_id", "label", "difficulty", "doi", "conditions_json")
RESPONSE_FIELDS = ("participant_id", "experience", "reaction_id", "response", "stored_score")


def wilson_interval(correct: int, count: int, z: float = 1.959963984540054):
    """Return a 95% Wilson interval, or missing bounds for an empty group."""
    if count < 0 or not 0 <= correct <= count:
        raise ValueError("Correct responses must be between zero and the response count.")
    if not count:
        return None, None
    p = correct / count
    denominator = 1 + z * z / count
    centre = (p + z * z / (2 * count)) / denominator
    half_width = z * sqrt(p * (1 - p) / count + z * z / (4 * count * count)) / denominator
    return max(0.0, centre - half_width), min(1.0, centre + half_width)


def _mean_interval(values):
    if not values:
        return None, None
    if len(values) < 2:
        return None, None
    # The workbook uses a t interval across participants on the fixed panel.
    from scipy.stats import t

    half_width = float(t.ppf(0.975, len(values) - 1)) * stdev(values) / sqrt(len(values))
    return max(0.0, mean(values) - half_width), min(1.0, mean(values) + half_width)


def _read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _write_csv(path, rows, fields=None):
    rows = list(rows)
    with Path(path).open("w", encoding="utf-8", newline="") as stream:
        columns = fields or (list(rows[0]) if rows else [])
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def _ratio(numerator, denominator):
    return numerator / denominator if denominator else None


def _summary(rows):
    count = len(rows)
    correct = sum(row["correct"] for row in rows)
    low, high = wilson_interval(correct, count)
    return {
        "responses": count, "correct": correct, "accuracy": _ratio(correct, count),
        "wilson_lower": low, "wilson_upper": high,
    }


def analyse(questions: list[dict], responses: list[dict]) -> dict[str, Any]:
    """Calculate accuracy, confidence summaries and binary agreement.

    Reaction IDs are the join key. Question numbers and response order have no
    role in scoring. Missing answers are excluded from accuracy denominators.
    Unknown reaction IDs and repeated participant/question pairs are errors.
    """
    panel = {}
    question_names = set()
    for question in questions:
        rid = question["reaction_id"]
        if not rid or rid in panel or question["question"] in question_names:
            raise ValueError("Question names and reaction IDs must be nonempty and unique.")
        if question["label"] not in ("P", "N"):
            raise ValueError(f"Invalid scoring label for {rid}.")
        panel[rid] = question
        question_names.add(question["question"])
    if not panel:
        raise ValueError("The question panel is empty.")

    by_participant = defaultdict(list)
    by_question = defaultdict(list)
    seen = set()
    for row in responses:
        pid, rid = row["participant_id"], row["reaction_id"]
        if not pid or rid not in panel:
            raise ValueError("Every response needs a participant ID and a reaction ID in the panel.")
        if (pid, rid) in seen:
            raise ValueError(f"Duplicate response for {pid}, {rid}.")
        seen.add((pid, rid))
        response = row["response"]
        if response not in RESPONSE_LABELS:
            raise ValueError(f"Unknown confidence category: {response!r}.")
        value = dict(row)
        value["predicted_label"] = RESPONSE_LABELS[response]
        value["correct"] = int(value["predicted_label"] == panel[rid]["label"])
        value["confidence"] = "Very Confident" if response.startswith("Very Confident") else "Likely"
        by_participant[pid].append(value)
        by_question[rid].append(value)

    participants = []
    for pid, rows in sorted(by_participant.items()):
        experiences = {row["experience"] for row in rows}
        stored_scores = {str(row.get("stored_score", "")) for row in rows}
        if len(experiences) != 1 or len(stored_scores) != 1:
            raise ValueError(f"Inconsistent participant metadata for {pid}.")
        stored_text = next(iter(stored_scores))
        stored = float(stored_text) if stored_text else None
        if stored is not None and (not 0 <= stored <= len(panel) or not stored.is_integer()):
            raise ValueError(f"Invalid stored score for {pid}.")
        correct = sum(row["correct"] for row in rows)
        participants.append({
            "participant_id": pid, "experience": next(iter(experiences)),
            "answered": len(rows), "missing": len(panel) - len(rows),
            "correct": correct, "accuracy": correct / len(rows),
            "stored_score": stored,
            "stored_minus_recalculated": stored - correct if stored is not None else None,
        })

    item_results = []
    for question in questions:
        rows = by_question[question["reaction_id"]]
        counts = Counter(row["response"] for row in rows)
        p_votes = sum(row["predicted_label"] == "P" for row in rows)
        n_votes = len(rows) - p_votes
        majority = "P" if p_votes > n_votes else "N" if n_votes > p_votes else "Tie" if rows else ""
        result = {key: question.get(key, "") for key in QUESTION_FIELDS if key != "conditions_json"}
        result.update(_summary(rows))
        result.update({
            "p_votes": p_votes, "n_votes": n_votes, "majority": majority,
            "majority_correct": int(majority == question["label"]) if majority in ("P", "N") else None,
            "consensus_share": _ratio(max(p_votes, n_votes), len(rows)),
            "pair_agreement": _ratio(p_votes * (p_votes - 1) + n_votes * (n_votes - 1), len(rows) * (len(rows) - 1)),
        })
        for category in RESPONSE_LABELS:
            result[category.lower().replace(" ", "_") + "_count"] = counts[category]
        item_results.append(result)

    all_rows = [row for rows in by_participant.values() for row in rows]
    experience_results = []
    experiences = list(EXPERIENCE_ORDER) + sorted({p["experience"] for p in participants} - set(EXPERIENCE_ORDER))
    for experience in experiences + ["Overall"]:
        group = participants if experience == "Overall" else [p for p in participants if p["experience"] == experience]
        scores = [p["correct"] for p in group]
        accuracies = [p["accuracy"] for p in group]
        low, high = _mean_interval(accuracies)
        experience_results.append({
            "experience": experience, "participants": len(group),
            "mean_score": mean(scores) if scores else None,
            "score_sd": stdev(scores) if len(scores) > 1 else None,
            "mean_accuracy": mean(accuracies) if accuracies else None,
            "t_lower": low, "t_upper": high,
        })

    confidence_results = []
    for experience in ["Overall"] + experiences:
        group = all_rows if experience == "Overall" else [r for r in all_rows if r["experience"] == experience]
        for category in list(RESPONSE_LABELS) + ["Very Confident", "Likely"]:
            rows = [row for row in group if row["response"] == category or row["confidence"] == category]
            confidence_results.append({"experience": experience, "confidence": category,
                                       **_summary(rows), "response_share": _ratio(len(rows), len(group))})

    valid_items = [item for item in item_results if item["responses"] >= 2]
    pair_agreement = mean(item["pair_agreement"] for item in valid_items) if valid_items else None
    p_fraction = mean(item["p_votes"] / item["responses"] for item in valid_items) if valid_items else None
    equal_ratings = len({item["responses"] for item in valid_items}) == 1
    chance = p_fraction ** 2 + (1 - p_fraction) ** 2 if p_fraction is not None else None
    fleiss = (pair_agreement - chance) / (1 - chance) if equal_ratings and chance is not None and chance < 1 else None
    ac1_chance = 2 * p_fraction * (1 - p_fraction) if p_fraction is not None else None
    ac1 = (pair_agreement - ac1_chance) / (1 - ac1_chance) if ac1_chance is not None else None
    rating_count = sum(item["responses"] for item in valid_items)
    p_total = sum(item["p_votes"] for item in valid_items)
    n_total = rating_count - p_total
    observed = _ratio(sum(2 * item["p_votes"] * item["n_votes"] / (item["responses"] - 1) for item in valid_items), rating_count)
    expected = _ratio(2 * p_total * n_total, rating_count * (rating_count - 1))
    alpha = 1 - observed / expected if expected else None
    majority_correct = [item["majority_correct"] for item in item_results if item["majority_correct"] is not None]
    label_accuracy = {
        label: _ratio(sum(item["correct"] for item in item_results if item["label"] == label),
                      sum(item["responses"] for item in item_results if item["label"] == label))
        for label in ("P", "N")
    }
    overall = _summary(all_rows)
    # Overall uncertainty is measured across participants on the fixed panel.
    overall.pop("wilson_lower")
    overall.pop("wilson_upper")
    overall_group = experience_results[-1]
    summary = {
        "participants": len(participants), "questions": len(panel), **overall,
        "mean_participant_accuracy": overall_group["mean_accuracy"],
        "participant_t_lower": overall_group["t_lower"],
        "participant_t_upper": overall_group["t_upper"],
        "success_prediction_fraction": _ratio(sum(item["p_votes"] for item in item_results), len(all_rows)),
        "positive_accuracy": label_accuracy["P"], "negative_accuracy": label_accuracy["N"],
        "balanced_accuracy": mean(label_accuracy.values()) if all(v is not None for v in label_accuracy.values()) else None,
        "majority_accuracy": mean(majority_correct) if majority_correct else None,
        "tied_questions": sum(item["majority"] == "Tie" for item in item_results),
        "raw_pairwise_agreement": pair_agreement, "fleiss_kappa": fleiss,
        "fleiss_equal_ratings": equal_ratings, "gwet_ac1": ac1, "krippendorff_alpha": alpha,
        "stored_score_mismatches": sum(p["stored_minus_recalculated"] not in (0, None) for p in participants),
    }
    return {"summary": summary, "participants": participants, "questions": item_results,
            "experience": experience_results, "confidence": confidence_results}


def load_benchmark(questions_file, responses_file):
    return analyse(_read_csv(questions_file), _read_csv(responses_file))


def compare_panels(human_questions, model_questions):
    """Compare the recorded human panel with a model-evaluation panel by ID."""
    human = {row["reaction_id"]: row for row in human_questions}
    model = {row["reaction_id"]: row for row in model_questions}
    if len(human) != len(human_questions) or len(model) != len(model_questions):
        raise ValueError("Each panel must contain unique reaction IDs.")
    differences = []
    for reaction_id in human.keys() & model.keys():
        human_row, model_row = human[reaction_id], model[reaction_id]
        left = json.loads(human_row["conditions_json"])
        right = model_row["conditions"]
        fields = {key: {"human": left.get(key), "model": right.get(key)}
                  for key in sorted(left.keys() | right.keys()) if left.get(key) != right.get(key)}
        if human_row["label"] != model_row["label"]:
            fields["label"] = {"human": human_row["label"], "model": model_row["label"]}
        if fields:
            differences.append({"question": human_row["question"], "reaction_id": reaction_id,
                                "differences": fields})
    differences.sort(key=lambda row: int(row["question"].removeprefix("Q")))
    return {"human_questions": len(human), "model_questions": len(model),
            "shared_reaction_ids": len(human.keys() & model.keys()),
            "human_only_ids": sorted(human.keys() - model.keys()),
            "model_only_ids": sorted(model.keys() - human.keys()),
            "all_shared_labels_match": all(human[key]["label"] == model[key]["label"] for key in human.keys() & model.keys()),
            "condition_differences": differences}


def write_analysis(result, output_dir):
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    (output / "summary.json").write_text(json.dumps(result["summary"], indent=2, allow_nan=False) + "\n", encoding="utf-8")
    for key in ("participants", "questions", "experience", "confidence"):
        _write_csv(output / f"{key}.csv", result[key])
    return output


def read_workbook(workbook_file):
    """Read the supplied workbook layout without exporting contact information.

    Sheet1 defines the panel. Sheet1.1 Responses contains a JSON object keyed by
    reaction ID in column C. Column D is a derived experience field, despite its
    heading.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_file, read_only=True, data_only=True)
    try:
        source = workbook["Sheet1"]
        questions = []
        for row in source.iter_rows(min_row=2, values_only=True):
            if not isinstance(row[14], (int, float)) or not row[16]:
                continue
            messages = json.loads(row[17])["messages"]
            condition_messages = [m["content"] for m in messages if m["role"] == "user"]
            labels = [m["content"] for m in messages if m["role"] == "assistant"]
            if len(condition_messages) != 1 or labels != [row[19]]:
                raise ValueError("Question JSON does not agree with the workbook scoring label.")
            questions.append({"question": f"Q{int(row[14])}", "reaction_id": row[16],
                              "label": row[19], "difficulty": row[18], "doi": row[15],
                              "conditions_json": json.dumps(json.loads(condition_messages[0]), ensure_ascii=False)})
        responses = []
        number = 0
        for row in workbook["Sheet1.1 Responses"].iter_rows(min_row=2, values_only=True):
            if row[2] is None:
                continue
            record = json.loads(row[2])
            number += 1
            experience = record.pop("User_Experience")
            for reaction_id, response in record.items():
                responses.append({"participant_id": f"P{number:03d}", "experience": experience,
                                  "reaction_id": reaction_id, "response": response,
                                  "stored_score": row[1]})
        result = analyse(questions, responses)
        reconciliation = _compare_workbook(workbook, result)
        return questions, responses, reconciliation
    finally:
        workbook.close()


def _compare_workbook(workbook, result):
    """Compare raw-answer calculations with the saved workbook values."""
    mismatches = []
    checked = 0

    def compare(sheet, address, calculated):
        nonlocal checked
        cached = workbook[sheet][address].value
        if cached is None or calculated is None:
            return
        checked += 1
        matches = abs(cached - calculated) < 1e-10 if isinstance(cached, (int, float)) and isinstance(calculated, (int, float)) else cached == calculated
        if not matches:
            mismatches.append({"sheet": sheet, "cell": address, "workbook_value": cached, "calculated": calculated})

    question_columns = {"E": "responses", "F": "correct", "G": "accuracy", "H": "wilson_lower",
                        "I": "wilson_upper", "J": "majority", "K": "majority_correct",
                        "L": "consensus_share", "M": "p_votes", "N": "n_votes", "O": "pair_agreement",
                        "U": "very_confident_success_count", "V": "likely_success_count",
                        "W": "likely_fail_count", "X": "very_confident_fail_count"}
    for index, question in enumerate(result["questions"], 6):
        for column, field in question_columns.items():
            compare("Sheet1.2 Questions", f"{column}{index}", question[field])
    for index, participant in enumerate(result["participants"], 8):
        for column, field in {"D": "stored_score", "E": "answered", "F": "correct", "G": "accuracy", "I": "stored_minus_recalculated"}.items():
            compare("Sheet1.3 Calculations", f"{column}{index}", participant[field])
    for index, group in enumerate(result["experience"], 11):
        for column, field in {"C": "participants", "D": "mean_score", "E": "score_sd", "F": "mean_accuracy", "G": "t_lower", "H": "t_upper"}.items():
            compare("Sheet2 Summary", f"{column}{index}", group[field])
    for index, confidence in enumerate(result["confidence"][:6], 41):
        for column, field in {"C": "responses", "D": "correct", "E": "accuracy"}.items():
            compare("Sheet2 Summary", f"{column}{index}", confidence[field])
    compare("Sheet2 Summary", "K6", result["summary"]["accuracy"])
    return {"checked_numeric_or_categorical_cells": checked, "saved_value_differences": mismatches,
            "stored_scores_matching_raw_answers": result["summary"]["participants"] - result["summary"]["stored_score_mismatches"]}


def export_workbook(workbook_file, output_dir):
    """Export deidentified scientific responses and the exact workbook panel."""
    questions, responses, reconciliation = read_workbook(workbook_file)
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    _write_csv(output / "human_questions.csv", questions, QUESTION_FIELDS)
    _write_csv(output / "human_responses.csv", responses, RESPONSE_FIELDS)
    manifest = {
        "source_id": "mof_quest_human_reference",
        "source_description": "Human MOF Quest question and response workbook",
        "source_sha256": hashlib.sha256(Path(workbook_file).read_bytes()).hexdigest(),
        "participants": len({row["participant_id"] for row in responses}),
        "questions": len(questions), "responses": len(responses),
        "participant_id_rule": "Sequential P001 identifiers in source response-row order; no contact fields exported.",
        "question_source": "Sheet1 columns O:T; response mapping uses the reaction ID in column Q.",
        "response_source": "Sheet1.1 Responses column C JSON; stored score from column B.",
        "response_labels": RESPONSE_LABELS,
        "files": {name: {"sha256": hashlib.sha256((output / name).read_bytes()).hexdigest()}
                  for name in ("human_questions.csv", "human_responses.csv")},
    }
    (output / "human_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    analysis = commands.add_parser("analyze", aliases=["analyse"], help="Analyze the anonymous benchmark CSV files.")
    analysis.add_argument("--questions", default="benchmarks/mof_quest/human_questions.csv")
    analysis.add_argument("--responses", default="benchmarks/mof_quest/human_responses.csv")
    analysis.add_argument("--output-dir", default="results/evaluation/human_quest")
    export = commands.add_parser("export-workbook", help="Export scientific fields from a local workbook.")
    export.add_argument("workbook")
    export.add_argument("--output-dir", required=True)
    comparison = commands.add_parser("compare-panels", help="Compare human and model question conditions by reaction ID.")
    comparison.add_argument("--human-questions", default="benchmarks/mof_quest/human_questions.csv")
    comparison.add_argument("--model-questions", default="benchmarks/mof_quest/questions.json")
    comparison.add_argument("--output", default="results/evaluation/human_quest/panel_comparison.json")
    args = parser.parse_args(argv)
    if args.command == "export-workbook":
        result = export_workbook(args.workbook, args.output_dir)
        print(json.dumps({key: result[key] for key in ("participants", "questions", "responses")}, indent=2))
    elif args.command == "compare-panels":
        result = compare_panels(_read_csv(args.human_questions), json.loads(Path(args.model_questions).read_text(encoding="utf-8")))
        output = Path(args.output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        result = load_benchmark(args.questions, args.responses)
        write_analysis(result, args.output_dir)
        print(json.dumps(result["summary"], indent=2, allow_nan=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
